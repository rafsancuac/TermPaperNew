#!/usr/bin/env python3
"""Check date formats, prefilled IDs, and dashboard formulas."""
import openpyxl

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(ROOT, "03_data_entry_template", "Marine_Fish_Marketing_Data_Entry (1).xlsx")
wb = openpyxl.load_workbook(PATH)

# Number formats of date/time cells
print("=== Number formats ===")
checks = [
    ("Form_A_Aratdar", "D5", "Interview_Date"),
    ("Form_C_Consumer", "D5", "Interview_Date"),
    ("Form_C_Consumer", "E5", "Interview_time"),
    ("Data_Collection_Log", "B5", "Date"),
    ("Data_Collection_Log", "F5", "Session_start"),
    ("Data_Collection_Log", "G5", "Session_end"),
    ("Form_M_Market_Observation", "D5", "Obs_Date"),
    ("Tag_Price_Sheet", "B5", "Date"),
]
for sheet, cell, label in checks:
    c = wb[sheet][cell]
    print(f"  {sheet}!{cell} ({label}): fmt='{c.number_format}' val={c.value!r}")

# Prefilled respondent IDs
print("\n=== Prefilled Respondent_IDs ===")
for name, col in [("Form_A_Aratdar", "B"), ("Form_B_Bepari_Faria", "B"),
                  ("Form_R_Khuchra", "B"), ("Form_C_Consumer", "B")]:
    ws = wb[name]
    ids = [ws[f"{col}{r}"].value for r in range(5, ws.max_row + 1)]
    ids = [i for i in ids if i]
    print(f"  {name}: {len(ids)} ids: {ids[0]} ... {ids[-1]} | first 12: {ids[:12]}")

# Price_Observations prefilled
ws = wb["Price_Observations"]
ids = [ws[f"B{r}"].value for r in range(5, ws.max_row + 1)]
ids = [i for i in ids if i]
print(f"  Price_Observations: {len(ids)} Obs_IDs: {ids[0]} ... {ids[-1]}")

ws = wb["Consumer_Purchases_Focal"]
print(f"  Consumer_Purchases_Focal max_row={ws.max_row} (rows 5+ empty template: {[ws[f'B{r}'].value for r in range(5,8)]})")

# Dashboard formulas
print("\n=== Progress_Dashboard formulas ===")
ws = wb["Progress_Dashboard"]
for addr in ["B6", "C6", "F6", "G6", "H6", "B12", "G12", "B15", "B17", "B18", "B21", "B23", "B24", "F6"]:
    print(f"  {addr}: {ws[addr].value}")

# Form_M 'done' check formula + column header check on QC
ws = wb["Progress_Dashboard"]
for r in range(6, 13):
    print(f"  row{r} F: {ws.cell(row=r, column=6).value}")

# Check formulas in a few Price_Obs cells deeper (row 6) to confirm relative refs
ws = wb["Price_Observations"]
print("\nPrice_Obs J6:", ws["J6"].value)
print("Price_Obs K6:", ws["K6"].value)
print("Price_Obs N6:", ws["N6"].value)

# Tag price sheet prefilled?
ws = wb["Tag_Price_Sheet"]
print("\nTag B5..B8:", [ws[f"B{r}"].value for r in range(5, 9)])

# Consumer sheets prefilled IDs?
for nm in ["Consumer_Purchases_Focal", "Consumer_Other_Fish"]:
    ws = wb[nm]
    print(f"{nm} B5:B8:", [ws[f"B{r}"].value for r in range(5, 9)])
