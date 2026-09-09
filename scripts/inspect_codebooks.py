#!/usr/bin/env python3
"""Dump full codebooks, instructions, and key validation lists."""
import openpyxl

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(ROOT, "03_data_entry_template", "Marine_Fish_Marketing_Data_Entry (1).xlsx")
wb = openpyxl.load_workbook(PATH, data_only=True)

# Full Instructions sheet
print("=== Instructions (full) ===")
ws = wb["Instructions"]
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=42, values_only=True), 1):
    cells = [str(c) if c is not None else "" for c in row]
    if any(cells):
        print(f"R{row_idx}: " + " | ".join(x for x in cells if x))

print("\n=== Codebook_Species (full) ===")
ws = wb["Codebook_Species"]
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=14, values_only=True), 1):
    cells = [str(c) if c is not None else "" for c in row]
    if any(cells):
        print(f"R{row_idx}: " + " | ".join(cells))

print("\n=== Codebook_Markets (full) ===")
ws = wb["Codebook_Markets"]
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
    cells = [str(c) if c is not None else "" for c in row]
    if any(cells):
        print(f"R{row_idx}: " + " | ".join(cells))

print("\n=== Unit_Converter (full) ===")
ws = wb["Unit_Converter"]
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=18, values_only=True), 1):
    cells = [str(c) if c is not None else "" for c in row]
    if any(cells):
        print(f"R{row_idx}: " + " | ".join(cells))

print("\n=== Progress_Dashboard (full) ===")
ws = wb["Progress_Dashboard"]
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=26, values_only=True), 1):
    cells = [str(c) if c is not None else "" for c in row]
    if any(cells):
        print(f"R{row_idx}: " + " | ".join(cells))

print("\n=== Data_Collection_Log rows 1-10 (formulas view) ===")
wb2 = openpyxl.load_workbook(PATH)  # formulas
ws = wb2["Data_Collection_Log"]
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
    cells = [str(c) if c is not None else "" for c in row]
    if any(cells):
        print(f"R{row_idx}: " + " | ".join(cells))
