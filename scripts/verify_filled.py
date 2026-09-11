#!/usr/bin/env python3
"""Verify the filled workbook: QC values, dashboard, chain consistency, validations.
Now supports REAL field data for Chattogram (04_data_filled/*REAL*.xlsx).
Simulated test data archived in 04_data_filled/archive/.
"""
import argparse
import openpyxl
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Use the same discovery logic as run_analysis.py
sys.path.insert(0, os.path.join(ROOT, "scripts"))
try:
    from run_analysis import find_default_data
    DEFAULT = find_default_data()
except Exception:
    DEFAULT = os.path.join(ROOT, "04_data_filled", "Marine_Fish_Marketing_Data_Entry_Chattogram_REAL_EMPTY.xlsx")

ap = argparse.ArgumentParser(description="Verify filled workbook (Chattogram real data)")
ap.add_argument("--data", default=DEFAULT, help="path to filled workbook (.xlsx)")
args = ap.parse_args()
OUT = args.data
print(f"Verifying workbook: {OUT}")

# ---------- cached values (post-recalc) ----------
try:
    wbv = openpyxl.load_workbook(OUT, data_only=True)
except Exception as e:
    print(f"[ERROR] Could not open {OUT}: {e}")
    print("If this is the EMPTY template, fill yellow cells with real field data first.")
    print("Simulated data is in 04_data_filled/archive/ for testing:")
    print("  python scripts/verify_filled.py --data 04_data_filled/archive/SIMULATED_v2_20260911_Chattogram_Filled.xlsx")
    sys.exit(1)

print("=== QC_Check (values) ===")
ws = wbv["QC_Check"]
# Sheet layout (header row 5): B=Check, C=Value, D=Rule / target, E=Status.
# (Bugfix 2026-09-09: status/rule were previously read from columns D/E swapped,
# so the script always reported 'non-passing checks: 24' even when every check
# passed. Status text is compared against the allowed vocabulary and the rule
# text is printed for context.)
fails = 0
total = 0
PASS_TOKENS = ("OK", "TARGET MET", "Info", "COMPLETE", "All matched", "No pairs yet")
for r in range(6, 30):
    label = ws.cell(row=r, column=2).value
    val = ws.cell(row=r, column=3).value
    rule = ws.cell(row=r, column=4).value
    status = ws.cell(row=r, column=5).value
    if label:
        total += 1
        ok = (status is None) or (str(status).strip() in PASS_TOKENS)
        print(f"  [{'PASS' if ok else 'FAIL'}] {str(label)[:52]:52s} = {val}"
              f"  [{str(status).strip() if status is not None else ''}]  rule: {rule}")
        if not ok:
            fails += 1
print(f"  >>> checks read: {total}; non-passing: {fails}")

print("\n=== Progress_Dashboard (values) ===")
ws = wbv["Progress_Dashboard"]
for r in range(5, 13):
    cells = [ws.cell(row=r, column=c).value for c in range(2, 9)]
    print("  ", [str(x)[:18] if x is not None else "" for x in cells])
for r in range(15, 25):
    label = ws.cell(row=r, column=2).value
    val = ws.cell(row=r, column=3).value
    if label:
        print(f"  {str(label)[:45]:45s} = {val}")

# ---------- formula workbook (validations survived?) ----------
wbf = openpyxl.load_workbook(OUT)
print("\n=== Data validation survival check ===")
for name in ["Form_A_Aratdar", "Form_B_Bepari_Faria", "Form_R_Khuchra", "Form_C_Consumer",
             "Price_Observations", "Consumer_Purchases_Focal", "Form_M_Market_Observation"]:
    n = len(wbf[name].data_validations.dataValidation)
    print(f"  {name}: {n} validation rules")

print("\n=== Blue formula columns intact (spot) ===")
ws = wbf["Price_Observations"]
print("  Price_Obs J5:", ws["J5"].value)
print("  Form_A J5 (cap kg):", wbf["Form_A_Aratdar"]["J5"].value)

# ---------- chain consistency spot checks ----------
print("\n=== Chain consistency (S01 Ilish @ M1, BDT/kg) ===")
ws = wbv["Price_Observations"]
rows = []
for r in range(5, 505):
    if ws.cell(row=r, column=6).value == "S01" and ws.cell(row=r, column=4).value == "M1":
        rows.append((r, ws.cell(row=r, column=3).value, ws.cell(row=r, column=5).value,
                     ws.cell(row=r, column=10).value, ws.cell(row=r, column=11).value))
for r in rows[:6]:
    print(f"  row{r[0]} {r[1]} {r[2]:13s} buy={r[3]} sell={r[4]}")

print("\n=== Consumer S01 prices @ M1 (should be near retail sell) ===")
ws = wbv["Consumer_Purchases_Focal"]
for r in range(5, 100):
    if ws.cell(row=r, column=4).value == "S01" and ws.cell(row=r, column=3).value == "M1":
        print(f"  row{r}: {ws.cell(row=r, column=2).value} paid={ws.cell(row=r, column=9).value}"
              f" qty={ws.cell(row=r, column=12).value}")

print("\n=== Pair price matching check ===")
ws = wbv["Price_Observations"]
pair_rows = {}
for r in range(5, 505):
    pid = ws.cell(row=r, column=15).value
    if pid:
        pair_rows.setdefault(pid, []).append(
            (ws.cell(row=r, column=3).value, ws.cell(row=r, column=6).value,
             ws.cell(row=r, column=10).value, ws.cell(row=r, column=11).value))
for pid, lst in list(pair_rows.items())[:5]:
    print(f"  {pid}: {lst}")

print("\n=== Payment sums (Form_A sample) ===")
ws = wbv["Form_A_Aratdar"]
bad = 0
for r in range(5, 41):
    y, z, aa = (ws.cell(row=r, column=c).value for c in (25, 26, 27))
    if y is not None and (y + z + aa) != 100:
        bad += 1
        print(f"  BAD row {r}: {y}+{z}+{aa}={y+z+aa}")
print(f"  Form_A payment rows not summing 100: {bad}")

print("\n=== Sample filled rows ===")
ws = wbv["Form_A_Aratdar"]
hdr = [ws.cell(row=4, column=c).value for c in range(2, 16)]
row5 = [ws.cell(row=5, column=c).value for c in range(2, 16)]
for h, v in zip(hdr, row5):
    print(f"    {str(h)[:34]:34s} = {v}")

print("\n=== K/D flag counts (Price_Obs) ===")
kbuy = ksell = dbuy = dsell = 0
for r in range(5, 505):
    g, h = ws.cell(row=r, column=7).value, ws.cell(row=r, column=8).value
    # re-open Price_Obs (ws was Form_A above)
ws = wbv["Price_Observations"]
for r in range(5, 505):
    g = ws.cell(row=r, column=7).value
    h = ws.cell(row=r, column=8).value
    kbuy += (g == "K"); dbuy += (g == "D")
    ksell += (h == "K"); dsell += (h == "D")
print(f"  buy: K={kbuy} D={dbuy} | sell: K={ksell} D={dsell}")

# Exit non-zero if any QC check failed (useful for CI/git hooks).
import sys as _sys
_sys.exit(1 if fails else 0)
