#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MS-499 - Marine Fish Marketing System, Chattogram (March 2026)
EXTENSION analyses (post-baseline), following scripts/run_analysis.py
conventions and Methodology V4 Sec. 3.8 / MASTER_PROMPT Part 3.

Run AFTER scripts/run_analysis.py:
    python scripts/run_analysis.py
    python scripts/extend_analysis.py

Outputs (analysis_outputs/):
  tables/T12_Pair_Detail.csv            matched-pair price rows (Pair_ID)
  tables/T12b_Wilcoxon_Result.csv       Wilcoxon signed-rank (consistency)
  tables/T13_Species_Market_Spread.csv  species x market price decomposition
  tables/T14_Species_Market_KruskalWallis.csv  per-species KW across markets
  tables/T14b_Dunn_Holm_Posthoc.csv     Dunn post-hoc (Holm), where applied
  tables/T15_Payment_Actor_ChiSquare.csv  payment mode x actor chi-square
  tables/T16_Stratum_Margin_MannWhitney.csv  stratum margin comparisons
  tables/T17_Retailer_MC_Profit_Spearman.csv  marketing cost vs profit
  charts/C8_retail_vs_wholesale_by_market.png
  and all tables appended to Analysis_Summary.xlsx (Index updated).

Nothing in 04_data_filled/ or the template is modified; results are derived.
"""
import argparse
import os
import sys

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from scipy import stats
from scipy.stats import rankdata, norm

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from run_analysis import (load_data, find_default_data, MAUND, LANDING_MARKETS,
                          TBL_DIR, CH_DIR, XLSX_OUT, HDR_FILL, BAND_FILL,
                          C_PRODUCER, C_ARATDAR, C_BEPARI, C_RETAIL, C_CONSUMER,
                          num)

SP_ORDER_C = None  # unused placeholder kept for clarity

plt.rcParams.update({
    "figure.dpi": 150, "savefig.dpi": 150,
    "font.size": 10, "axes.titlesize": 12, "axes.titleweight": "bold",
    "axes.labelsize": 10, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.alpha": 0.25, "grid.linestyle": "--",
    "legend.frameon": False,
})

EXT_SHEETS = ["T12 Pair Detail", "T12b Wilcoxon", "T13 Species-Market Spread",
              "T14 Kruskal-Wallis", "T14b Dunn-Holm", "T15 Payment x Actor",
              "T16 Stratum Margins", "T17 MC-Prof Spearman"]


# ----------------------------------------------------------------------------
# 1. Matched pairs (Methodology 3.3.4) & Wilcoxon signed-rank
# ----------------------------------------------------------------------------
# Channel order: an Aratdar sells to a Bepari/Faria or retailer; Bepari/Faria
# sells to retailers; a retailer sells to consumers. In a matched buyer-seller
# pair the SELLER's row must be at the same or an upstream stage of the buyer.
ACTOR_ORDER = {"Aratdar": 0, "Bepari_Faria": 1, "Khuchra": 2}


def build_pairs(po):
    """Pair rows into matched transactions. Each pair carries the seller's
    row and the buyer's row; we compare the seller's sell quote with the
    buyer's buy quote of the same lot. Alignment respects channel order so
    that a missing buyer-side quote marks the pair incomplete rather than
    silently pairing a seller's sell with the buyer's own (different) buy."""
    raw = {}
    for r in po:
        pid = r.get("Pair_ID")
        if not pid:
            continue
        raw.setdefault(pid, []).append(r)

    out = []
    for pid in sorted(raw):
        rows = raw[pid]
        if len(rows) < 2:
            continue
        best = None  # (distance, seller_row, buyer_row)
        for a in range(len(rows)):
            for b in range(len(rows)):
                if a == b:
                    continue
                s, bb = rows[a], rows[b]
                if s["sell_kg"] is None or bb["buy_kg"] is None:
                    continue
                if ACTOR_ORDER.get(s.get("Actor_type"), 9) > \
                        ACTOR_ORDER.get(bb.get("Actor_type"), 9):
                    continue  # buyer cannot be upstream of the seller
                d = abs(s["sell_kg"] - bb["buy_kg"])
                if best is None or d < best[0]:
                    best = (d, s, bb)
        if best is None:
            out.append({"Pair_ID": pid, "Market": rows[0].get("Market"),
                        "Species_code": rows[0].get("Species_code"),
                        "Seller_ID": None, "Seller_actor": None,
                        "Seller_sell_BDT_kg": None,
                        "Buyer_ID": None, "Buyer_actor": None,
                        "Buyer_buy_BDT_kg": None, "Diff_BDT_kg": None,
                        "Complete": False})
            continue
        _, sa, bb = best
        out.append({"Pair_ID": pid,
                    "Market": sa.get("Market") or bb.get("Market"),
                    "Species_code": sa.get("Species_code"),
                    "Seller_ID": sa.get("Respondent_ID"),
                    "Seller_actor": sa.get("Actor_type"),
                    "Seller_sell_BDT_kg": round(sa["sell_kg"], 2),
                    "Buyer_ID": bb.get("Respondent_ID"),
                    "Buyer_actor": bb.get("Actor_type"),
                    "Buyer_buy_BDT_kg": round(bb["buy_kg"], 2),
                    "Diff_BDT_kg": round(sa["sell_kg"] - bb["buy_kg"], 2),
                    "Complete": True})
    return pd.DataFrame(out)


def table_pair_detail(po):
    df = build_pairs(po)
    note = ("Matched-pair transaction prices (Methodology 3.3.4). Seller_sell = the "
            "seller's quoted sale price of the lot; Buyer_buy = the buyer's quoted "
            "purchase price of the same lot. Diff = seller sell - buyer buy (≈ 0 if "
            "self-reports are consistent). Complete = both sides quoted numerically.")
    return "T12_Pair_Detail", "Table 12. Matched buyer-seller pairs (Pair_ID)", df, note


def table_wilcoxon(po):
    df = build_pairs(po)
    comp = df[df["Complete"] & df["Diff_BDT_kg"].notna()]
    diffs = comp["Diff_BDT_kg"].to_numpy()
    n = len(diffs)
    rows = []
    if n == 0:
        rows.append({"Item": "Wilcoxon signed-rank (two-sided)",
                     "Value": "No usable pairs"})
    else:
        try:
            res = stats.wilcoxon(diffs, zero_method="wilcox",
                                 alternative="two-sided", correction=False)
            w, p = float(res.statistic), float(res.pvalue)
        except ValueError:
            # all-zero differences cannot be tested
            w, p = None, None
        rows.append({"Item": "Usable matched pairs (n)", "Value": n})
        rows.append({"Item": "Pairs with non-zero difference", "Value": int(np.sum(diffs != 0))})
        rows.append({"Item": "Median difference (BDT/kg)", "Value": round(float(np.median(diffs)), 2)})
        rows.append({"Item": "Mean difference (BDT/kg)", "Value": round(float(np.mean(diffs)), 2)})
        rows.append({"Item": "Wilcoxon W statistic", "Value": w})
        rows.append({"Item": "p-value (two-sided)", "Value": round(p, 4) if p is not None else None})
        rows.append({"Item": "Interpretation", "Value": (
            "p >= 0.05: seller-sell and buyer-buy quotes of matched lots do not differ "
            "systematically - self-reported chain prices are internally consistent"
            if (p is not None and p >= 0.05)
            else "p < 0.05: matched buy/sell reports differ systematically - flag for review")})
    dfout = pd.DataFrame(rows)
    note = ("Wilcoxon signed-rank on paired differences (seller's sell - buyer's buy) of the "
            "same matched lot (Methodology 3.3.4). Diagnostic consistency check, not a powered "
            "hypothesis test; report n alongside the result.")
    return "T12b_Wilcoxon_Result", "Table 12b. Wilcoxon signed-rank test on matched pairs", dfout, note


# ----------------------------------------------------------------------------
# 2. Species x market decomposition & Kruskal-Wallis (Methodology 3.8)
# ----------------------------------------------------------------------------
def _series_by(po, cpf_yes, code, actor, field, market):
    vals = [r[field] for r in po
            if r.get("Species_code") == code and r.get("Market") == market
            and r.get("Actor_type") == actor and r[field] is not None]
    return vals


def table_species_market_spread(d):
    po = d["PO"]
    cpf_yes = [r for r in d["CPF"] if r.get("Purchased_today") == "Yes"]
    mkts = ["M1", "M2", "M3", "M4", "M5", "M6"]
    rows = []
    for code in sorted(d["SP"]):
        for m in mkts:
            ws = _series_by(po, cpf_yes, code, "Bepari_Faria", "sell_kg", m)
            rs = _series_by(po, cpf_yes, code, "Khuchra", "sell_kg", m)
            cs = [r["price_kg"] for r in cpf_yes
                  if r.get("Species_code") == code and r.get("Market") == m
                  and r["price_kg"] is not None]
            wm = round(float(np.mean(ws)), 2) if ws else None
            rm = round(float(np.mean(rs)), 2) if rs else None
            cm = round(float(np.mean(cs)), 2) if cs else None
            # quote-based (vendor) retail margin and consumer-paid (realised) margin
            m_q = round(rm - wm, 2) if (rm is not None and wm is not None) else None
            m_c = round(cm - wm, 2) if (cm is not None and wm is not None) else None
            rows.append({"Species": code,
                         "Market": m,
                         "Wholesale_sell_mean": wm, "n_wholesale": len(ws),
                         "Retailer_sell_quote_mean": rm, "n_retailer": len(rs),
                         "Consumer_paid_mean": cm, "n_consumer": len(cs),
                         "Retail_margin_quote_BDT_kg": m_q,
                         "Retail_margin_paid_BDT_kg": m_c})
    df = pd.DataFrame(rows)
    note = ("Species x market price decomposition. Wholesale = Bepari/Faria sell quotes; "
            "Retailer quote = khuchra sell quotes; Consumer = Form C slips. The difference "
            "between quote- and paid-based margins reflects bargaining at retail. Cells with "
            "n=0 are not observed. Per Methodology 3.8 cells with n<5 support descriptive "
            "statements only.")
    return "T13_Species_Market_Spread", "Table 13. Species x market price spread (BDT/kg)", df, note


def _kruskal_by_market(vals_by_market, min_cell):
    """vals_by_market: {market: [prices]}. KW over markets with >= min_cell obs.
    Returns stats dict or None if fewer than 3 qualifying markets."""
    qual = {m: v for m, v in vals_by_market.items() if len(v) >= min_cell}
    if len(qual) < 3:
        return None, qual
    groups = [np.asarray(v, dtype=float) for v in qual.values()]
    h, p = stats.kruskal(*groups)
    return {"qualifying_markets": sorted(qual),
            "n_per_market": {m: len(v) for m, v in qual.items()},
            "H": round(float(h), 3), "p": float(p), "df": len(qual) - 1}, qual


def dunn_holm(groups, names, alpha=0.05):
    """Dunn's post-hoc with Holm-Bonferroni step-down (Methodology 3.8)."""
    pooled = np.concatenate(groups)
    n = len(pooled)
    ranks = rankdata(pooled)
    # tie correction T = sum(t^3 - t)
    _, counts = np.unique(pooled, return_counts=True)
    tie_corr = float(np.sum(counts ** 3 - counts) / (12 * (n - 1)))
    var_rank = n * (n + 1) / 12 - tie_corr
    out = []
    idx = np.cumsum([0] + [len(g) for g in groups])
    grp_means = [float(np.mean(ranks[idx[i]:idx[i + 1]])) for i in range(len(groups))]
    pairs = [(i, j) for i in range(len(groups)) for j in range(i + 1, len(groups))]
    raw_p = []
    for i, j in pairs:
        se = np.sqrt(var_rank * (1.0 / len(groups[i]) + 1.0 / len(groups[j])))
        z = (grp_means[i] - grp_means[j]) / se
        p = 2 * (1 - norm.cdf(abs(z)))
        raw_p.append((i, j, z, p))
    # Holm-Bonferroni step-down: adj_p(k) = min(1, p(k)*(m-k+1)),
    # then enforce monotonicity adj_p non-decreasing in k.
    order = sorted(range(len(raw_p)), key=lambda k: raw_p[k][3])
    adj_prev = 0.0
    for rank_k in range(len(order)):
        i, j, z, p = raw_p[order[rank_k]]
        adj = min(1.0, p * (len(order) - rank_k))
        adj = max(adj, adj_prev)
        adj_prev = adj
        sig = adj < alpha
        out.append({"Market_i": names[i], "Market_j": names[j],
                    "z": round(float(z), 3), "raw_p": round(p, 4),
                    "p_holm": round(adj, 4), "significant_0.05": sig})
    return out


def table_species_market_kw(d):
    po = d["PO"]
    mkts = ["M1", "M2", "M3", "M4", "M5", "M6"]
    rows, dunn_rows = [], []
    for code in sorted(d["SP"]):
        vbm = {m: _series_by(po, None, code, "Khuchra", "sell_kg", m) for m in mkts}
        resA, qualA = _kruskal_by_market(vbm, 5)
        resB, qualB = _kruskal_by_market(vbm, 3)
        # prefer the pre-specified cell size (n>=5); else exploratory n>=3
        chosen, tier, res = resA, "pre-specified (n>=5)", resA
        if res is None and resB is not None:
            chosen, tier, res = resB, "exploratory (n>=3)", resB
        if res is None:
            rows.append({"Species": code, "Local_name": d["SP"][code]["Local_name"],
                         "Test_tier": "not run", "Markets_qualifying": "",
                         "n_per_market": "", "H": None, "df": None, "p_value": None,
                         "Decision": "cells < n threshold in <3 markets -> descriptive only"})
            continue
        cells = {m: len(v) for m, v in vbm.items() if len(v) >= 1}
        rows.append({"Species": code, "Local_name": d["SP"][code]["Local_name"],
                     "Test_tier": tier,
                     "Markets_qualifying": "; ".join(res["qualifying_markets"]),
                     "n_per_market": "; ".join(f"{m}={res['n_per_market'][m]}"
                                               for m in res["qualifying_markets"]),
                     "H": res["H"], "df": res["df"], "p_value": round(res["p"], 4),
                     "Decision": "reject H0 (p<0.05): prices differ across markets"
                     if res["p"] < 0.05 else "fail to reject H0 (p>=0.05)"})
        if res["p"] < 0.05:
            groups = [np.asarray(vbm[m], dtype=float) for m in res["qualifying_markets"]]
            for dr in dunn_holm(groups, res["qualifying_markets"]):
                dr = dict(dr)
                dr.update({"Species": code})
                dunn_rows.append(dr)
    df = pd.DataFrame(rows)
    note = ("Kruskal-Wallis H (Methodology 3.8): do retailer selling prices of the same "
            "species differ across markets? Cells with n>=5 in >=3 markets use the "
            "pre-specified test; otherwise an exploratory tier (n>=3) is flagged and species "
            "with thinner coverage are reported descriptively only (no pooled-across-species "
            "tests).")
    return "T14_Species_Market_KruskalWallis", "Table 14. Kruskal-Wallis: species price across markets", df, note


def table_dunn(d):
    # recompute to fill Dunn rows (mirror of the previous table builder)
    po = d["PO"]
    mkts = ["M1", "M2", "M3", "M4", "M5", "M6"]
    rows = []
    for code in sorted(d["SP"]):
        vbm = {m: _series_by(po, None, code, "Khuchra", "sell_kg", m) for m in mkts}
        resA, _ = _kruskal_by_market(vbm, 5)
        resB, _ = _kruskal_by_market(vbm, 3)
        res = resA if resA is not None else resB
        if res is not None and res["p"] < 0.05:
            groups = [np.asarray(vbm[m], dtype=float) for m in res["qualifying_markets"]]
            for dr in dunn_holm(groups, res["qualifying_markets"]):
                dr = dict(dr)
                dr.update({"Species": code})
                rows.append(dr)
    df = pd.DataFrame(rows)
    note = "Dunn's post-hoc with Holm step-down (p_holm) following significant Kruskal-Wallis."
    return "T14b_Dunn_Holm_Posthoc", "Table 14b. Dunn-Holm pairwise market comparisons", df, note


# ----------------------------------------------------------------------------
# 3. Payment mode x actor (chi-square) - Methodology 3.8 Q4
# ----------------------------------------------------------------------------
def table_payment_actor(d):
    freq_col = "MFS_transaction_frequency"
    cats = ["Regular", "Occasional", "Never"]
    counts = {}
    for label, rs in (("Aratdar", d["A"]), ("Bepari/Faria", d["B"]), ("Retailer", d["R"])):
        counts[label] = [sum(1 for r in rs if r.get(freq_col) == c) for c in cats]
    arr = np.array([counts[k] for k in counts], dtype=int)
    # expected counts
    chi2, p, dof, exp = stats.chi2_contingency(arr)
    n = arr.sum()
    cramer = float(np.sqrt(chi2 / (n * (min(arr.shape) - 1))))
    sparse = exp.min() < 5
    g_text = ""
    if sparse:
        _, g_p, _, _ = stats.chi2_contingency(arr, lambda_="log-likelihood")
        g_text = f", G-test p={g_p:.4f}"
    rows = []
    for i, actor in enumerate(counts):
        row = {"Actor": actor, "n": int(arr[i].sum()),
               "MFS_Regular": int(arr[i][0]), "MFS_Occasional": int(arr[i][1]),
               "MFS_Never": int(arr[i][2]),
               "MFS_share_pct": round(100.0 * arr[i][0] / arr[i].sum(), 1)}
        rows.append(row)
    ptext = f"chi2={chi2:.2f}, df={dof}, p={p:.4f}{g_text}"
    rows.append({"Actor": "chi-square (actor x MFS frequency)",
                 "n": int(n), "MFS_Regular": "", "MFS_Occasional": "",
                 "MFS_Never": "", "MFS_share_pct": ptext})
    rows.append({"Actor": "Cramer's V / smallest expected count",
                 "n": "", "MFS_Regular": "", "MFS_Occasional": "",
                 "MFS_Never": "", "MFS_share_pct": f"V={cramer:.3f} / min exp={exp.min():.2f}"})
    df = pd.DataFrame(rows)
    note = ("Payment-mode adoption by trader class (Methodology 3.8 Q4). Chi-square on "
            "MFS-use frequency (Regular/Occasional/Never); where expected cell counts < 5 the "
            "likelihood-ratio (G) test is reported alongside Pearson chi-square as the sparse-"
            "cell fallback. Consumer payment method is a separate question and appears in "
            "Tables 5/8.")
    return "T15_Payment_Actor_ChiSquare", "Table 15. Payment mode x actor chi-square", df, note


# ----------------------------------------------------------------------------
# 4. Stratum margins (Mann-Whitney U) - Methodology 3.8 Q2
# ----------------------------------------------------------------------------
def respondent_margins(po):
    """Per respondent: mean own-quote margin (sell-buy, rows with both prices)."""
    acc = {}
    for r in po:
        if r["buy_kg"] is None or r["sell_kg"] is None:
            continue
        key = (r.get("Respondent_ID"), r.get("Actor_type"), r.get("Market"))
        acc.setdefault(key, []).append(r["sell_kg"] - r["buy_kg"])
    return {k: float(np.mean(v)) for k, v in acc.items()}


def table_stratum_margins(d):
    po = d["PO"]
    margins = respondent_margins(po)
    strata = {"Aratdar": [], "Bepari_Faria": [], "Khuchra": []}
    for (rid, actor, mkt), m in margins.items():
        if actor in strata:
            strata[actor].append(m)
    rows = []
    for actor, vals in strata.items():
        vals = np.asarray(vals)
        rows.append({"Actor": actor, "n": len(vals),
                     "Median_margin_BDT_kg": round(float(np.median(vals)), 2),
                     "IQR_BDT_kg": round(float(np.percentile(vals, 75) - np.percentile(vals, 25)), 2),
                     "Mean_margin_BDT_kg": round(float(np.mean(vals)), 2)})
    pairs = [("Aratdar", "Bepari_Faria"), ("Aratdar", "Khuchra"), ("Bepari_Faria", "Khuchra")]
    for a, b in pairs:
        u, p = stats.mannwhitneyu(strata[a], strata[b], alternative="two-sided")
        ptext = f"U={u:.0f}, p<0.0001" if p < 0.0001 else f"U={u:.0f}, p={p:.4f}"
        rows.append({"Actor": f"Mann-Whitney U: {a} vs {b}",
                     "n": f"{len(strata[a])}/{len(strata[b])}",
                     "Median_margin_BDT_kg": "", "IQR_BDT_kg": "",
                     "Mean_margin_BDT_kg": ptext})
    df = pd.DataFrame(rows)
    note = ("Per-respondent mean own-quote margin = mean of (sell - buy) over species "
            "quoted with both prices on the interview day (Methodology 3.7.2: M = P_s - P_b). "
            "Mann-Whitney U per stratum pair (Q2); Holm adjustment is recommended when "
            "quoting all three comparisons.")
    return "T16_Stratum_Margin_MannWhitney", "Table 16. Trader stratum margins and Mann-Whitney tests", df, note


# ----------------------------------------------------------------------------
# 5. Marketing cost vs net margin (Spearman) - retailers (Methodology 3.8 Q5)
# ----------------------------------------------------------------------------
def table_mc_profit_spearman(d):
    po = d["PO"]
    margins = respondent_margins(po)
    rows = []
    for r in d["R"]:
        rid = r.get("Respondent_ID")
        key = next((k for k in margins if k[0] == rid), None)
        if key is None:
            continue
        cap = num(r.get("Daily_capacity_kg")) or num(r.get("Daily_capacity_raw")) or np.nan
        if not np.isfinite(cap) or cap <= 0:
            continue
        mc = ((num(r.get("Shop_Van_Rent_BDT_per_day")) or 0)
              + (num(r.get("Ice_cost_BDT_per_day")) or 0)
              + (num(r.get("Wash_Water_Other_cost_BDT_per_day")) or 0)) / cap
        pi = margins[key] - mc
        rows.append({"Respondent_ID": rid, "Market": r.get("Market"),
                     "Mean_margin_BDT_kg": round(margins[key], 2),
                     "MC_BDT_kg": round(float(mc), 2),
                     "Net_profit_BDT_kg": round(float(pi), 2)})
    df = pd.DataFrame(rows)
    out = []
    if len(df) >= 8:
        rho, p = stats.spearmanr(df["MC_BDT_kg"], df["Net_profit_BDT_kg"])
        out = df.copy()
        summary = pd.DataFrame([{"Respondent_ID": "Spearman rho(MC, net margin)",
                                 "Market": f"n={len(df)}", "Mean_margin_BDT_kg": "",
                                 "MC_BDT_kg": round(float(rho), 3),
                                 "Net_profit_BDT_kg": round(float(p), 4)}])
        out = pd.concat([df, summary], ignore_index=True)
    note = ("Retailers only (daily cost structure is unambiguous): MC/kg = (stall rent + ice + "
            "wash/water/other, all BDT/day) / daily capacity kg; net margin = own-quote margin "
            "- MC/kg. Spearman rho between MC and net margin (Methodology 3.8 Q5). Other strata "
            "mix monthly/yearly/per-trip frequencies, so per-kg MC is not computed for them "
            "until the cost module defines the period consistently (flagged in review).")
    return "T17_Retailer_MC_Profit_Spearman", "Table 17. Retailer marketing cost vs net margin (Spearman)", out, note


# ----------------------------------------------------------------------------
# Chart C8
# ----------------------------------------------------------------------------
def chart_c8(d, species_with_data=None):
    po = d["PO"]
    mkts = ["M1", "M2", "M3", "M4", "M5", "M6"]
    panels = []
    for code in sorted(d["SP"]):
        nm = sum(1 for m in mkts
                 if (len(_series_by(po, None, code, "Bepari_Faria", "sell_kg", m)) >= 1
                     and len(_series_by(po, None, code, "Khuchra", "sell_kg", m)) >= 1))
        if nm >= 4:
            panels.append(code)
    panels = panels[:8]
    if not panels:
        return None
    nrow = int(np.ceil(len(panels) / 2))
    fig, axes = plt.subplots(nrow, 2, figsize=(12.5, 3.6 * nrow),
                             constrained_layout=True)
    axes = np.atleast_1d(axes).ravel()
    for ax, code in zip(axes, panels):
        wm = [_series_by(po, None, code, "Bepari_Faria", "sell_kg", m) for m in mkts]
        rm = [_series_by(po, None, code, "Khuchra", "sell_kg", m) for m in mkts]
        ws = [float(np.mean(v)) if v else np.nan for v in wm]
        rs = [float(np.mean(v)) if v else np.nan for v in rm]
        x = np.arange(len(mkts))
        ax.bar(x - 0.18, ws, 0.36, color=C_BEPARI, label="Wholesale (bepari sell)", alpha=0.9)
        ax.bar(x + 0.18, rs, 0.36, color=C_RETAIL, label="Retail (khuchra sell)", alpha=0.9)
        for xi, (w, r) in enumerate(zip(ws, rs)):
            nw = len(wm[xi]); nr = len(rm[xi])
            for vv, n in ((w, nw), (r, nr)):
                if not np.isnan(vv):
                    ax.text(xi + (0.18 if vv == r else -0.18), vv + 8,
                            str(n), ha="center", fontsize=7, color="#555555")
        ax.set_xticks(x)
        ax.set_xticklabels([d["MK"][m] if len(d["MK"][m]) <= 12 else d["MK"][m][:11] + "."
                            for m in mkts], rotation=25, ha="right", fontsize=7.5)
        ax.set_title(f"{code} {d['SP'][code]['Local_name']}")
        ax.set_ylim(0, max([v for v in ws + rs if not np.isnan(v)] or [1]) * 1.15)
        ax.grid(axis="x", visible=False)
    for ax in axes[len(panels):]:
        ax.axis("off")
    axes[0].legend(ncol=2, loc="upper left", fontsize=8.5)
    fig.suptitle("Wholesale vs retail price by market, Chattogram, March 2026 (labels = n quotes)",
                 fontsize=13, fontweight="bold")
    path = os.path.join(CH_DIR, "C8_retail_vs_wholesale_by_market.png")
    fig.savefig(path)
    plt.close(fig)
    return path


# ----------------------------------------------------------------------------
# Writer: append extension tables to Analysis_Summary.xlsx
# ----------------------------------------------------------------------------
def append_xlsx(tables):
    wb = openpyxl.load_workbook(XLSX_OUT)
    for sheet, title, df, note in tables:
        name = sheet[:31]
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        ws.cell(row=2, column=2, value=title).font = Font(bold=True, size=12,
                                                           color=HDR_FILL)
        ws.cell(row=3, column=2, value=note).font = Font(italic=True, size=9,
                                                          color="808080")
        headers = list(df.columns)
        for j, h in enumerate(headers):
            c = ws.cell(row=4, column=2 + j, value=str(h))
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor=HDR_FILL)
        for i, (_, row) in enumerate(df.iterrows()):
            for j, h in enumerate(headers):
                v = row[h]
                if pd.isna(v):
                    v = None
                elif isinstance(v, (np.integer,)):
                    v = int(v)
                elif isinstance(v, (np.floating,)):
                    v = float(v)
                c = ws.cell(row=5 + i, column=2 + j, value=v)
                if isinstance(v, float):
                    c.number_format = "0.00" if abs(v) < 100 else "0.0"
                if i % 2 == 1:
                    c.fill = PatternFill("solid", fgColor=BAND_FILL)
    # Index maintenance
    idx = wb["Index"]
    headers = ["Sheet", "Table title", "Rows", "Cols"]
    j0 = 2
    present = {idx.cell(row=r, column=2).value for r in range(5, idx.max_row + 1)}
    next_row = 5
    while idx.cell(row=next_row, column=2).value is not None:
        next_row += 1
    for sheet, title, df, note in tables:
        if sheet[:31] in present:
            continue
        for j, v in enumerate((sheet[:31], title, len(df), len(df.columns))):
            idx.cell(row=next_row, column=j0 + j, value=v)
        next_row += 1
    wb.save(XLSX_OUT)
    print(f"  appended {len(tables)} extension sheets to {os.path.basename(XLSX_OUT)}")


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="MS-499 extension analysis")
    ap.add_argument("--data", default=None)
    args = ap.parse_args()
    path = args.data or find_default_data()
    d = load_data(path)
    print(f"Reading workbook: {path}")
    print(f"Extension analyses on loaded data: A={len(d['A'])}, B={len(d['B'])}, "
          f"R={len(d['R'])}, C={len(d['C'])}, price_obs={len(d['PO'])}")

    tables = [table_pair_detail(d["PO"]), table_wilcoxon(d["PO"]),
              table_species_market_spread(d), table_species_market_kw(d),
              table_dunn(d), table_payment_actor(d), table_stratum_margins(d),
              table_mc_profit_spearman(d)]
    os.makedirs(TBL_DIR, exist_ok=True)
    os.makedirs(CH_DIR, exist_ok=True)
    for sheet, title, df, note in tables:
        csv = os.path.join(TBL_DIR, f"{sheet}.csv")
        df.to_csv(csv, index=False, encoding="utf-8-sig")
        print(f"  table  {sheet}.csv  ({len(df)} rows)")
    ch = chart_c8(d)
    if ch:
        print(f"  chart  {os.path.basename(ch)}")
    else:
        print("  chart  C8 skipped: <4 markets with wholesale+retail quotes for every species")
    append_xlsx(tables)
    print("DONE - extension outputs in analysis_outputs/")


if __name__ == "__main__":
    main()
