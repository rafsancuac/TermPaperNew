#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MS-499 Term Paper - Marine Fish Marketing System, Chattogram (March 2026)
Analysis pipeline for the filled survey workbook.

Input : ../04_data_filled/Marine_Fish_Marketing_Data_Entry_Chattogram_Filled.xlsx
        (override with:  python run_analysis.py --data <path-to-filled-xlsx>)
Output: ../analysis_outputs/
          tables/*.csv          - every analysis table as CSV (utf-8-sig)
          charts/*.png          - 7 publication-ready figures (English labels)
          Analysis_Summary.xlsx - all tables in one styled workbook

Portable: needs only Python 3.9+, openpyxl, pandas, matplotlib, numpy.
Re-run after replacing simulated values with real field data - everything
recomputes automatically.

Note on language: chart/table labels are in English because the survey
workbook entries were designed in English; Bengali commentary for each output
is provided in the repository README and docs/WRITING_GUIDE.md.
"""
import argparse
import os
from collections import Counter

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
MAUND = 37.32  # 1 maund = 37.32 kg (template Unit_Converter)
# First-sale (fisher->aratdar auction) prices are observable only at the two
# landing-linked markets (Methodology V4 Table 3.1 & Sec. 3.11c): Fishery Ghat
# and Patenga. Producer-side and auction (aratdar sell) series are therefore
# restricted to these markets; wholesale/retail series cover all six markets.
LANDING_MARKETS = ("M1", "M6")

TBL_DIR = os.path.join(ROOT, "analysis_outputs", "tables")
CH_DIR = os.path.join(ROOT, "analysis_outputs", "charts")
XLSX_OUT = os.path.join(ROOT, "analysis_outputs", "Analysis_Summary.xlsx")

# Self-contained palette (kept identical to the project design system)
C_PRODUCER, C_ARATDAR, C_BEPARI, C_RETAIL, C_CONSUMER = (
    "#1f4e79", "#d9822b", "#2e8b57", "#c0504d", "#7b5ea7")
HDR_FILL = "1F4E79"
BAND_FILL = "F2F6FA"

EDU_ORDER = ["Illiterate", "Primary", "Secondary", "Higher_Secondary"]
MFS_ORDER = ["Regular", "Occasional", "Never"]

plt.rcParams.update({
    "figure.dpi": 150, "savefig.dpi": 150,
    "font.size": 10, "axes.titlesize": 12.5, "axes.titleweight": "bold",
    "axes.labelsize": 10.5, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.alpha": 0.25, "grid.linestyle": "--",
    "legend.frameon": False,
})


# ----------------------------------------------------------------------------
# Generic helpers
# ----------------------------------------------------------------------------
def num(v):
    """Return float for numeric cells, else None (K/D flags, blanks)."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def mean_sd(vals, nd=2):
    v = [x for x in vals if x is not None]
    if not v:
        return None, None, 0
    sd = round(float(np.std(v, ddof=1)), nd) if len(v) > 1 else 0.0
    return round(float(np.mean(v)), nd), sd, len(v)


def price_per_kg(raw, unit, cached):
    """BDT/kg price with fallback recomputation from the raw quote."""
    v = num(cached)
    if v is not None:
        return round(v, 2)
    v = num(raw)
    if v is None:
        return None
    return round(v / MAUND if str(unit or "").strip() == "Maund" else v, 2)


def qty_to_kg(raw, unit, cached):
    """Quantity in kg with fallback recomputation (maund -> kg multiplies)."""
    v = num(cached)
    if v is not None:
        return round(v, 2)
    v = num(raw)
    if v is None:
        return None
    return round(v * MAUND if str(unit or "").strip() == "Maund" else v, 2)


def read_sheet_rows(ws, id_col=2, require=None):
    """Read a template sheet (headers on row 4, data from row 5)."""
    headers = {}
    for c in range(2, ws.max_column + 1):
        h = ws.cell(row=4, column=c).value
        if h is not None and str(h).strip():
            headers[c] = str(h).strip()
    rows = []
    for r in range(5, ws.max_row + 1):
        rid = ws.cell(row=r, column=id_col).value
        if rid in (None, ""):
            continue
        rec = {h: ws.cell(row=r, column=c).value for c, h in headers.items()}
        if require and any(rec.get(k) in (None, "") for k in require):
            continue
        rows.append(rec)
    return rows


def find_default_data():
    d = os.path.join(ROOT, "04_data_filled")
    if os.path.isdir(d):
        for f in sorted(os.listdir(d)):
            if f.lower().endswith(".xlsx") and not f.startswith("~$"):
                return os.path.join(d, f)
    raise FileNotFoundError("No filled workbook found in 04_data_filled/")


# ----------------------------------------------------------------------------
# Data loading
# ----------------------------------------------------------------------------
def load_data(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    d = {"path": path}
    d["species"] = read_sheet_rows(wb["Codebook_Species"])
    d["markets"] = read_sheet_rows(wb["Codebook_Markets"])
    d["SP"] = {r["Code"]: r for r in d["species"]}
    d["MK"] = {r["Code"]: r["Market_name"] for r in d["markets"]}

    d["A"] = read_sheet_rows(wb["Form_A_Aratdar"], require=["Interview_Date"])
    d["B"] = read_sheet_rows(wb["Form_B_Bepari_Faria"], require=["Interview_Date"])
    d["R"] = read_sheet_rows(wb["Form_R_Khuchra"], require=["Interview_Date"])
    d["C"] = read_sheet_rows(wb["Form_C_Consumer"], require=["Interview_Date"])

    po = read_sheet_rows(wb["Price_Observations"], require=["Respondent_ID"])
    for r in po:
        r["buy_kg"] = price_per_kg(r.get("Buy_price_raw"), r.get("Unit"),
                                    r.get("Buy_BDT_per_kg"))
        r["sell_kg"] = price_per_kg(r.get("Sell_price_raw"), r.get("Unit"),
                                     r.get("Sell_BDT_per_kg"))
        r["qty_kg"] = qty_to_kg(r.get("Quantity_raw"), r.get("Quantity_unit"),
                                 r.get("Quantity_kg"))
    d["PO"] = po

    cpf = read_sheet_rows(wb["Consumer_Purchases_Focal"])
    for r in cpf:
        r["price_kg"] = price_per_kg(r.get("Price_raw"), r.get("Unit"),
                                      r.get("Price_BDT_per_kg"))
        r["qty_kg"] = qty_to_kg(r.get("Quantity_raw"), r.get("Quantity_unit"),
                                 r.get("Quantity_kg"))
    d["CPF"] = cpf

    cof = read_sheet_rows(wb["Consumer_Other_Fish"], require=["Fish_name_local"])
    for r in cof:
        r["price_kg"] = price_per_kg(r.get("Price_raw"), r.get("Unit"),
                                      r.get("Price_BDT_per_kg"))
    d["COF"] = cof

    d["M"] = read_sheet_rows(wb["Form_M_Market_Observation"], require=["Obs_Date"])
    d["TAG"] = read_sheet_rows(wb["Tag_Price_Sheet"], require=["Market"])

    dates = []
    for key in ("A", "B", "R", "C"):
        for r in d[key]:
            if r.get("Interview_Date") is not None:
                dates.append(pd.Timestamp(r["Interview_Date"]).date())
    d["window"] = (min(dates).strftime("%d/%m/%Y"), max(dates).strftime("%d/%m/%Y"))
    return d


# ----------------------------------------------------------------------------
# Table builders   (each returns DataFrame; collected in build_tables)
# ----------------------------------------------------------------------------
def _pct(counter, key, n):
    return round(100.0 * counter.get(key, 0) / n, 1) if n else 0.0


def t1_profile(d):
    groups = [
        ("Aratdar", d["A"]),
        ("Bepari", [r for r in d["B"] if r.get("Actor_subtype") == "Bepari"]),
        ("Faria", [r for r in d["B"] if r.get("Actor_subtype") == "Faria"]),
        ("Retailer", d["R"]),
    ]
    rows = []
    all_tr = []
    for label, rs in groups:
        all_tr += rs
        rows.append(_profile_row(label, rs))
    rows.append(_profile_row("All traders", all_tr))
    df = pd.DataFrame(rows)
    note = (f"Traders n={len(all_tr)}; consumers n={len(d['C'])} "
            f"(profile items not collected for consumers). "
            f"Survey window {d['window'][0]}-{d['window'][1]}.")
    return "T1_Respondent_Profile", "Table 1. Socio-economic profile of respondents", df, note


def _profile_row(label, rs):
    n = len(rs)
    age_m, age_sd, _ = mean_sd([num(r.get("Age_years")) for r in rs])
    yr_m, yr_sd, _ = mean_sd([num(r.get("Years_in_business")) for r in rs])
    fam_m, _, _ = mean_sd([num(r.get("Family_members")) for r in rs])
    edu = Counter(r.get("Education") for r in rs)
    row = {"Actor": label, "n": n,
           "Age_mean_years": age_m, "Age_sd_years": age_sd,
           "Years_in_business_mean": yr_m, "Years_in_business_sd": yr_sd,
           "Family_size_mean": fam_m}
    for e in EDU_ORDER:
        row[f"Edu_{e.replace('_', ' ')}_pct"] = _pct(edu, e, n)
    return row


def t2_scale(d):
    rows = []
    for label, rs, key in (("Aratdar", d["A"], "Daily_capacity_kg"),
                           ("Bepari/Faria", d["B"], "Daily_capacity_kg"),
                           ("Retailer", d["R"], "Daily_capacity_kg")):
        vals = [num(r.get(key)) for r in rs]
        m, sd, n = mean_sd(vals)
        v = [x for x in vals if x is not None]
        row = {"Actor": label, "n": n, "Daily_capacity_kg_mean": m,
               "Daily_capacity_kg_sd": sd,
               "Daily_capacity_kg_median": round(float(np.median(v)), 1) if v else None,
               "Daily_capacity_kg_min": round(min(v), 1) if v else None,
               "Daily_capacity_kg_max": round(max(v), 1) if v else None}
        if label == "Aratdar":
            row["Workers_mean"] = mean_sd([num(r.get("Workers_count")) for r in rs])[0]
            row["Worker_wage_BDT_day_mean"] = mean_sd(
                [num(r.get("Worker_daily_wage_BDT")) for r in rs])[0]
        if label == "Bepari/Faria":
            row["Trips_per_period_mean"] = mean_sd(
                [num(r.get("Trips_per_period")) for r in rs])[0]
            row["Ice_use_kg_day_mean"] = mean_sd(
                [num(r.get("Ice_usage_kg_per_day")) for r in rs])[0]
        if label == "Retailer":
            row["Spoilage_pct_mean"] = mean_sd(
                [num(r.get("Spoilage_pct")) for r in rs])[0]
        rows.append(row)
    df = pd.DataFrame(rows)
    note = "Daily traded/handled volume per actor group (kg/day)."
    return "T2_Business_Scale", "Table 2. Business scale of trading actors", df, note


def t2b_channel(d):
    rows = []

    def block(form_label, rs, field, values=None):
        cnt = Counter(r.get(field) for r in rs)
        keys = values if values else [k for k, _ in cnt.most_common()]
        for k in keys:
            if k in (None, ""):
                continue
            rows.append({"Form": form_label, "Field": field.replace("_", " "),
                         "Value": str(k), "n": cnt[k],
                         "pct": _pct(cnt, k, len(rs))})

    bep = d["B"]
    block("Form B (Bepari/Faria)", bep, "Actor_subtype")
    block("Form B (Bepari/Faria)", bep, "Business_pattern")
    block("Form B (Bepari/Faria)", bep, "Supplier_category")
    block("Form B (Bepari/Faria)", bep, "Source_location")
    block("Form B (Bepari/Faria)", bep, "Credit_sales_share",
          ["Almost_all", "Some", "None"])
    block("Form R (Retailer)", d["R"], "Buys_from")
    block("Form R (Retailer)", d["R"], "Sells_to")
    block("Form R (Retailer)", d["R"], "Sells_on_credit",
          ["Almost_all", "Some", "None"])
    df = pd.DataFrame(rows)
    note = "Channel structure: who supplies whom and on what terms (counts, % of form respondents)."
    return "T2b_Channel_Patterns", "Table 2b. Channel structure and trading patterns", df, note


def _po_mean(po, actor, field, species=None, markets=None):
    """Mean of numeric price/quantity values for actor+species(+markets)."""
    vals = [r[field] for r in po
            if r[field] is not None and r.get("Actor_type") == actor
            and (species is None or r.get("Species_code") == species)
            and (markets is None or r.get("Market") in markets)]
    return round(float(np.mean(vals)), 2) if vals else None


def _cons_mean(cpf, species=None):
    """Mean consumer-paid price (Form C focal purchases, 'Yes' today)."""
    vals = [r["price_kg"] for r in cpf
            if r.get("Purchased_today") == "Yes" and r["price_kg"] is not None
            and (species is None or r.get("Species_code") == species)]
    return round(float(np.mean(vals)), 2) if vals else None


def t3_chain(d):
    """Price chain per species.

    Methodological basis (Methodology V4, Sec. 3.7):
    - Producer (fisher first-sale) price = Form A BUY quotes at the two
      landing-linked markets only (M1 Fishery Ghat, M6 Patenga), where an
      aratdar's buy price is the net auction price paid to the fisherman.
    - Aratdar (auction hammer) sell = Form A SELL quotes at M1/M6 only.
    - Bepari sell = Form B sell quotes pooled over markets (city wholesale).
    - Retailer sell quote = Form R sell quotes pooled over markets
      (descriptive vendor quote, not used for margins).
    - Consumer price = mean actually paid on Form C focal purchases
      (consumer slips validate the retail price, Sec. 3.7.3).
    - Segment margins are differences of consecutive level means, so on any
      species with all levels present the margins telescope to the total
      spread (consumer - producer) exactly.
    - Species without a full chain (all four chain levels observed) are
      shown descriptively with blank margins; the ALL row and Table 10 pool
      only chain-complete species (state this in the Methods chapter).
    """
    po = d["PO"]
    cpf_yes = [r for r in d["CPF"] if r.get("Purchased_today") == "Yes"]
    rows = []
    for code in sorted(d["SP"].keys()):
        sp = d["SP"][code]
        n_obs = sum(1 for r in po if r.get("Species_code") == code
                    and (r.get("buy_kg") is not None or r.get("sell_kg") is not None))
        prod = _po_mean(po, "Aratdar", "buy_kg", code, LANDING_MARKETS)
        arat = _po_mean(po, "Aratdar", "sell_kg", code, LANDING_MARKETS)
        bep = _po_mean(po, "Bepari_Faria", "sell_kg", code)
        ret = _po_mean(po, "Khuchra", "sell_kg", code)
        cons = _cons_mean(cpf_yes, code)
        complete = all(x is not None for x in (prod, arat, bep, cons))

        def diff(a, b):
            return round(a - b, 2) if (a is not None and b is not None) else None

        row = {
            "Species_code": code, "Local_name": sp["Local_name"],
            "English_name": sp["English_common_name"], "n_price_obs": n_obs,
            "Producer_BDT_kg": prod, "Aratdar_sell_BDT_kg": arat,
            "Bepari_sell_BDT_kg": bep, "Retailer_sell_BDT_kg": ret,
            "Consumer_paid_BDT_kg": cons,
            "Aratdar_margin_BDT_kg": diff(arat, prod) if complete else None,
            "Bepari_margin_BDT_kg": diff(bep, arat) if complete else None,
            "Retailer_margin_BDT_kg": diff(cons, bep) if complete else None,
            "Total_spread_BDT_kg": diff(cons, prod) if complete else None,
            "Producer_share_pct": (round(100.0 * prod / cons, 1)
                                   if (prod is not None and cons) else None),
        }
        rows.append(row)

    # pooled row over chain-complete species only: unweighted mean of the
    # per-species means (composition-consistent; matches Table 10), so the
    # pooled margins telescope to the pooled spread and PS + spread % = 100.
    def pooled(field, complete_only=False):
        vals = [row[field] for row in rows if row[field] is not None
                and (not complete_only
                     or row["Total_spread_BDT_kg"] is not None)]
        return round(float(np.mean(vals)), 2) if vals else None

    full = [row for row in rows if row["Total_spread_BDT_kg"] is not None]
    codes_all = [row["Species_code"] for row in full]
    n_all = sum(1 for r in po if r.get("Species_code") in codes_all
                and (r.get("buy_kg") is not None or r.get("sell_kg") is not None))
    prod_p, arat_p = pooled("Producer_BDT_kg", True), pooled("Aratdar_sell_BDT_kg", True)
    bep_p, ret_p = pooled("Bepari_sell_BDT_kg", True), pooled("Retailer_sell_BDT_kg", True)
    cons_p = pooled("Consumer_paid_BDT_kg", True)
    rows.append({
        "Species_code": "ALL", "Local_name": f"Mean of species means (chain-complete: {', '.join(codes_all)})",
        "English_name": "-", "n_price_obs": n_all,
        "Producer_BDT_kg": prod_p, "Aratdar_sell_BDT_kg": arat_p,
        "Bepari_sell_BDT_kg": bep_p, "Retailer_sell_BDT_kg": ret_p,
        "Consumer_paid_BDT_kg": cons_p,
        "Aratdar_margin_BDT_kg": round(arat_p - prod_p, 2),
        "Bepari_margin_BDT_kg": round(bep_p - arat_p, 2),
        "Retailer_margin_BDT_kg": round(cons_p - bep_p, 2),
        "Total_spread_BDT_kg": round(cons_p - prod_p, 2),
        "Producer_share_pct": round(100.0 * prod_p / cons_p, 1),
    })
    df = pd.DataFrame(rows)
    note = ("Producer price = net auction price paid to fishers (Form A buy) at the "
            "landing-linked markets M1/M6 only; Aratdar sell = auction hammer price at M1/M6; "
            "Bepari sell = Form B sell over all markets; Retailer sell = Form R vendor quote "
            "(descriptive; margins use the consumer-paid anchor); Consumer = focal-species "
            "purchases actually paid (Form C). Margins = differences of consecutive level means "
            "and telescope to the total spread. K/D-flagged prices excluded. Species without a "
            "complete chain show blank margins and are excluded from ALL (see Local_name note).")
    return "T3_Price_Chain", "Table 3. Price chain by species (BDT per kg)", df, note


def t4_costs(d):
    rows = []

    def add(actor, item, unit, rs, field, note_txt=""):
        m, sd, n = mean_sd([num(r.get(field)) for r in rs])
        rows.append({"Actor": actor, "Cost_item": item, "Unit": unit,
                     "n": n, "Mean": m, "SD": sd, "Note": note_txt})

    A, B, Rp = d["A"], d["B"], d["R"]
    add("Aratdar", "Commission rate", "% of lot value",
        [r for r in A if r.get("Commission_basis") == "Percent"],
        "Commission_rate", "Percent-basis aratdars only")
    add("Aratdar", "Arat/office rent (yearly payers)", "BDT per year",
        [r for r in A if r.get("Rent_frequency") == "Yearly"], "Rent_amount_BDT")
    add("Aratdar", "Arat/office rent (monthly payers)", "BDT per month",
        [r for r in A if r.get("Rent_frequency") == "Monthly"], "Rent_amount_BDT")
    add("Aratdar", "Electricity bill (monthly payers)", "BDT per month",
        [r for r in A if r.get("Electricity_frequency") == "Monthly"],
        "Electricity_bill_BDT")
    add("Aratdar", "Electricity bill (yearly payers)", "BDT per year",
        [r for r in A if r.get("Electricity_frequency") == "Yearly"],
        "Electricity_bill_BDT")
    add("Aratdar", "Workers employed", "persons",
        A, "Workers_count", "Own/family + hired")
    add("Aratdar", "Worker daily wage", "BDT per day", A, "Worker_daily_wage_BDT")
    add("Aratdar", "Equipment & cooler maintenance (yearly)", "BDT per year",
        [r for r in A if r.get("Equipment_Cooler_Maintenance_frequency") == "Yearly"],
        "Equipment_Cooler_Maintenance_BDT")
    add("Aratdar", "Other operating cost", "BDT per period", A, "Other_cost_BDT",
        "See Other_cost_description in Form A")
    add("Bepari/Faria", "Transport fare per trip", "BDT per trip", B,
        "Transport_fare_per_trip_BDT")
    add("Bepari/Faria", "Trips per period", "trips", B, "Trips_per_period")
    add("Bepari/Faria", "Loading/unloading cost", "BDT per trip", B,
        "Loading_Unloading_cost_per_trip_BDT")
    add("Bepari/Faria", "Ice usage", "kg per day", B, "Ice_usage_kg_per_day")
    add("Bepari/Faria", "Ice cost", "BDT per day", B, "Ice_cost_BDT_per_day")
    add("Bepari/Faria", "Commission paid to aratdar (percent basis)", "% of lot",
        [r for r in B if r.get("Commission_unit") == "Percent_of_lot"],
        "Commission_to_aratdar_value")
    add("Bepari/Faria", "Commission paid to aratdar (flat basis)", "BDT per lot",
        [r for r in B if r.get("Commission_unit") == "BDT_per_lot"],
        "Commission_to_aratdar_value")
    add("Bepari/Faria", "Transit damage/loss", "% of lot", B, "Transit_damage_pct")
    add("Retailer", "Shop/van rent", "BDT per day", Rp, "Shop_Van_Rent_BDT_per_day")
    add("Retailer", "Ice cost", "BDT per day", Rp, "Ice_cost_BDT_per_day")
    add("Retailer", "Washing, water & other", "BDT per day", Rp,
        "Wash_Water_Other_cost_BDT_per_day")
    add("Retailer", "Spoilage/unsold loss", "% of purchase", Rp, "Spoilage_pct")
    df = pd.DataFrame(rows)
    note = "Mean marketing costs by actor (mixed frequencies kept separate - do not sum rows)."
    return "T4_Marketing_Costs", "Table 4. Marketing costs by channel actor", df, note


def t5_payment(d):
    base_keys = {"Cash_pct_mean": None, "MFS_pct_mean": None,
                 "Credit_pct_mean": None, "MFS_Regular_n": None,
                 "MFS_Occasional_n": None, "MFS_Never_n": None,
                 "Payment_method": None, "Payment_method_pct": None}
    rows = []
    for label, rs in (("Aratdar", d["A"]), ("Bepari/Faria", d["B"]),
                      ("Retailer", d["R"])):
        n = len(rs)
        cnt = Counter(r.get("MFS_transaction_frequency") for r in rs)
        row = dict(base_keys)
        row.update({
            "Actor": label, "n": n,
            "Cash_pct_mean": mean_sd([num(r.get("Payment_cash_pct")) for r in rs])[0],
            "MFS_pct_mean": mean_sd([num(r.get("Payment_MFS_pct")) for r in rs])[0],
            "Credit_pct_mean": mean_sd([num(r.get("Payment_credit_pct")) for r in rs])[0],
            "MFS_Regular_n": cnt.get("Regular", 0),
            "MFS_Occasional_n": cnt.get("Occasional", 0),
            "MFS_Never_n": cnt.get("Never", 0),
        })
        rows.append(row)

    # consumer payment method rows (Form C)
    cnt = Counter(r.get("Payment_method") for r in d["C"])
    for k, n_ in cnt.most_common():
        row = dict(base_keys)
        row.update({"Actor": "Consumer (Form C)", "n": len(d["C"]),
                    "Payment_method": str(k),
                    "Payment_method_pct": round(100.0 * n_ / len(d["C"]), 1)})
        rows.append(row)
    df = pd.DataFrame(rows)
    note = ("Trader rows: mean share of sales receipts settled by cash / mobile financial "
            "services (bKash, Nagad) / credit. Consumer row: % choosing each payment method.")
    return "T5_Payment_Methods", "Table 5. Payment methods across the channel", df, note


def t6_problems(d):
    probs = set()
    for rs, keys in ((d["A"], ("Problem_1", "Problem_2", "Problem_3")),
                     (d["B"], ("Problem_1", "Problem_2", "Problem_3")),
                     (d["R"], ("Problem_1", "Problem_2", "Problem_3"))):
        for r in rs:
            for k in keys:
                if r.get(k):
                    probs.add(r[k])
    rows = []
    nA, nB, nR = len(d["A"]), len(d["B"]), len(d["R"])

    def mentions(rs, p):
        return sum(1 for r in rs for k in ("Problem_1", "Problem_2", "Problem_3")
                   if r.get(k) == p)

    for p in probs:
        a, b, rr = mentions(d["A"], p), mentions(d["B"], p), mentions(d["R"], p)
        tot = a + b + rr
        rows.append({"Problem": p,
                     "Aratdar_n": a, "Aratdar_pct": round(100.0 * a / nA, 1),
                     "Bepari_n": b, "Bepari_pct": round(100.0 * b / nB, 1),
                     "Retailer_n": rr, "Retailer_pct": round(100.0 * rr / nR, 1),
                     "Total_n": tot, "Total_pct": round(100.0 * tot / (nA + nB + nR), 1)})
    rows.sort(key=lambda x: -x["Total_n"])
    df = pd.DataFrame(rows)
    note = "% = share of that actor's respondents mentioning the problem in any of the 3 slots (multi-response)."
    return "T6_Problems", "Table 6. Marketing problems reported by traders", df, note


def t7_market(d):
    rows = []
    for r in d["M"]:
        rows.append({
            "Market": r.get("Market"), "Market_name": r.get("Market_name"),
            "Obs_date": pd.Timestamp(r.get("Obs_Date")).strftime("%d/%m/%Y"),
            "Type": r.get("Market_type"), "Trading_hours": r.get("Trading_hours"),
            "Retail_stalls": r.get("Retail_stall_count"),
            "Transport_access": r.get("Transport_access"),
            "Platform": r.get("Platform_condition"), "Roofing": r.get("Roofing"),
            "Drainage": r.get("Drainage"), "Electricity": r.get("Electricity"),
            "Ice": r.get("Ice_availability"), "Sanitation": r.get("Sanitation"),
            "Water_supply": r.get("Water_supply"),
            "Fee_arrangement": r.get("Fee_arrangement"),
            "Fee_rate_note": r.get("Fee_rate_note"),
            "GPS": r.get("GPS_coordinates"),
        })
    df = pd.DataFrame(rows)
    note = "Form M direct market observation, one row per market."
    return "T7_Market_Infrastructure", "Table 7. Market infrastructure (Form M observation)", df, note


def t8_consumer(d):
    C = d["C"]
    CPF = d["CPF"]
    rows = []

    def block(field, label):
        cnt = Counter(r.get(field) for r in C)
        for k, n_ in cnt.most_common():
            if k in (None, ""):
                continue
            rows.append({"Item": label, "Value": str(k), "n": n_,
                         "pct": round(100.0 * n_ / len(C), 1)})

    block("Visit_frequency", "Visit frequency")
    block("Payment_method", "Payment method")
    block("Main_reason_choosing_arat", "Main reason for choosing seller")
    m, _, n = mean_sd([num(r.get("Visit_interval_days")) for r in C])
    rows.append({"Item": "Mean visit interval (days)", "Value": m, "n": n, "pct": None})
    cnt = Counter(r.get("Bought_from") for r in CPF)
    for k, n_ in cnt.most_common():
        if k:
            rows.append({"Item": "Focal purchase source (per purchase)", "Value": str(k),
                         "n": n_, "pct": round(100.0 * n_ / len(CPF), 1)})
    yes = [r for r in CPF if r.get("Purchased_today") == "Yes"]
    m, _, n = mean_sd([r["price_kg"] for r in yes])
    rows.append({"Item": "Mean focal-species price paid (BDT/kg)", "Value": m,
                 "n": n, "pct": None})
    m, _, n = mean_sd([r["qty_kg"] for r in yes])
    rows.append({"Item": "Mean purchase size (kg)", "Value": m, "n": n, "pct": None})
    df = pd.DataFrame(rows)
    note = f"Consumers n={len(C)}; focal purchases n={len(CPF)} (Yes today n={len(yes)})."
    return "T8_Consumer_Behaviour", "Table 8. Consumer behaviour", df, note


def t9_other_fish(d):
    rows = []
    cnt = Counter(r.get("Fish_name_local") for r in d["COF"])
    for name, n_ in cnt.most_common():
        rs = [r for r in d["COF"] if r.get("Fish_name_local") == name]
        rows.append({"Fish (local name)": name, "n": n_,
                     "Mean_price_BDT_kg": mean_sd([r.get("price_kg") for r in rs])[0],
                     "Mean_qty_kg": mean_sd([num(r.get("Quantity_raw")) for r in rs])[0]})
    df = pd.DataFrame(rows)
    note = "Non-focal marine items consumers reported buying (Consumer_Other_Fish sheet)."
    return "T9_Other_Fish", "Table 9. Other fish/items purchased by consumers", df, note


def t10_margins(d):
    """Channel margins & producer's share (composition-consistent pooling over
    chain-complete species only - see t3_chain docstring)."""
    po = d["PO"]
    cpf_yes = [r for r in d["CPF"] if r.get("Purchased_today") == "Yes"]
    codes = [c for c in sorted(d["SP"].keys())]

    # chain-complete species set (same rule as Table 3)
    chain_codes = []
    for c in codes:
        if (_po_mean(po, "Aratdar", "buy_kg", c, LANDING_MARKETS) is not None
                and _po_mean(po, "Aratdar", "sell_kg", c, LANDING_MARKETS) is not None
                and _po_mean(po, "Bepari_Faria", "sell_kg", c) is not None
                and _cons_mean(cpf_yes, c) is not None):
            chain_codes.append(c)
    if not chain_codes:
        chain_codes = codes

    def sp_mean(actor, field, mkts=None, cset=None):
        vals = []
        for c in (cset or chain_codes):
            m = _po_mean(po, actor, field, c, mkts)
            if m is not None:
                vals.append(m)
        return round(float(np.mean(vals)), 2) if vals else None, len(vals)

    prod, n_prod = sp_mean("Aratdar", "buy_kg", LANDING_MARKETS)
    arat, n_arat = sp_mean("Aratdar", "sell_kg", LANDING_MARKETS)
    bep, n_bep = sp_mean("Bepari_Faria", "sell_kg")
    cons_vals = []
    for c in chain_codes:
        v = _cons_mean(cpf_yes, c)
        if v is not None:
            cons_vals.append(v)
    cons = round(float(np.mean(cons_vals)), 2) if cons_vals else None

    def level(name, sell, buy, n_sp):
        if sell is None or buy is None:
            return {"Level": name, "n_species": n_sp, "Margin_BDT_kg": None,
                    "Margin_pct_consumer": None}
        mg = round(sell - buy, 2)
        return {"Level": name, "n_species": n_sp,
                "Margin_BDT_kg": mg,
                "Margin_pct_consumer": round(100.0 * mg / cons, 1) if cons else None}

    rows = [level("Aratdar (auction margin, M1/M6)", arat, prod, n_arat),
            level("Bepari/Faria (wholesale margin)", bep, arat, n_bep),
            level("Retailer (margin to consumer)", cons, bep, len(chain_codes))]
    rows.append({"Level": "Total marketing spread (consumer - producer)",
                 "n_species": len(chain_codes),
                 "Margin_BDT_kg": round(cons - prod, 2) if cons and prod else None,
                 "Margin_pct_consumer": round(100.0 - 100.0 * prod / cons, 1)
                 if cons and prod else None})
    rows.append({"Level": "Producer share of consumer price",
                 "n_species": len(chain_codes),
                 "Margin_BDT_kg": None,
                 "Margin_pct_consumer": round(100.0 * prod / cons, 1)
                 if cons and prod else None})
    df = pd.DataFrame(rows)
    note = ("Unweighted means of species-level means over chain-complete species only "
            f"(n_species={len(chain_codes)}; see Table 3 note). Producer & auction prices from "
            "landing-linked markets M1/M6 (first-sale proxy, Methodology 3.11c); wholesale "
            "price = Bepari sell over all markets; retail price = consumer-paid anchor (Form C "
            "slips). Margins are differences of consecutive level means, so they sum exactly to "
            "the total spread and PS% + spread% = 100. Margin % is expressed as a share of the "
            "pooled consumer price.")
    return "T10_Margin_Summary", "Table 10. Channel margins and producer's share", df, note


def t11_price_by_market(d):
    po = d["PO"]
    mkts = ["M1", "M2", "M3", "M4", "M5", "M6"]
    rows = []
    for code in sorted(d["SP"].keys()):
        sp = d["SP"][code]
        row = {"Species": f"{code} {sp['Local_name']}"}
        for m in mkts:
            vals = [r["sell_kg"] for r in po
                    if r.get("Actor_type") == "Khuchra" and r.get("Market") == m
                    and r.get("Species_code") == code and r["sell_kg"] is not None]
            row[d["MK"][m]] = round(float(np.mean(vals)), 0) if vals else None
        rows.append(row)
    df = pd.DataFrame(rows)
    note = "Mean retailer (khuchra) selling price, BDT/kg, by market. Blank = not observed."
    return "T11_Retail_Price_by_Market", "Table 11. Retail prices by species and market", df, note


def build_tables(d):
    return [t1_profile(d), t2_scale(d), t2b_channel(d), t3_chain(d),
            t4_costs(d), t5_payment(d), t6_problems(d), t7_market(d),
            t8_consumer(d), t9_other_fish(d), t10_margins(d), t11_price_by_market(d)]


# ----------------------------------------------------------------------------
# Charts
# ----------------------------------------------------------------------------
def make_charts(d, tables):
    tmap = {t[0]: t[2] for t in tables}

    # C1 - price chain by species -------------------------------------------
    df = tmap["T3_Price_Chain"]
    df = df[df["Species_code"] != "ALL"].sort_values(
        "Consumer_paid_BDT_kg", ascending=False, na_position="last").reset_index(drop=True)
    labels = [f"{r.Species_code}\n{r.Local_name}" for r in df.itertuples()]
    x = np.arange(len(df))
    w = 0.16
    series = [
        ("Producer (net auction)", "Producer_BDT_kg", C_PRODUCER),
        ("Aratdar sell", "Aratdar_sell_BDT_kg", C_ARATDAR),
        ("Bepari sell", "Bepari_sell_BDT_kg", C_BEPARI),
        ("Retailer sell", "Retailer_sell_BDT_kg", C_RETAIL),
        ("Consumer paid", "Consumer_paid_BDT_kg", C_CONSUMER),
    ]
    fig, ax = plt.subplots(figsize=(12.5, 6), constrained_layout=True)
    for i, (lab, col, colr) in enumerate(series):
        ax.bar(x + (i - 2) * w, df[col], w, label=lab, color=colr, edgecolor="white",
               linewidth=0.4)
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel("BDT per kg")
    ax.set_title("Marine fish price chain by species - Chattogram markets, March 2026")
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), ncol=5, borderaxespad=0)
    ax.grid(axis="y")
    ax.grid(axis="x", visible=False)
    fig.savefig(os.path.join(CH_DIR, "C1_price_chain_by_species.png"))
    plt.close(fig)

    # C2 - producer's share --------------------------------------------------
    df = tmap["T3_Price_Chain"]
    df = df[df["Species_code"] != "ALL"].dropna(
        subset=["Producer_share_pct"]).sort_values("Producer_share_pct")
    overall = tmap["T3_Price_Chain"].loc[
        tmap["T3_Price_Chain"]["Species_code"] == "ALL", "Producer_share_pct"].iloc[0]
    fig, ax = plt.subplots(figsize=(9, 6), constrained_layout=True)
    ax.barh(df["Species_code"] + " " + df["Local_name"], df["Producer_share_pct"],
            color=C_PRODUCER, alpha=0.85)
    ax.axvline(overall, color=C_CONSUMER, linestyle="--", linewidth=1.6)
    ax.text(overall + 0.8, 0.2, f"Pooled mean {overall:.0f}%", color=C_CONSUMER,
            fontsize=9.5, va="bottom")
    for i, v in enumerate(df["Producer_share_pct"]):
        ax.text(v - 0.8, i, f"{v:.0f}", ha="right", va="center", fontsize=9,
                color="white")
    ax.set_xlabel("Producer share of consumer price (%)")
    ax.set_xlim(0, max(df["Producer_share_pct"]) * 1.18)
    ax.set_title("Producer's share of the consumer's taka, by species")
    ax.grid(axis="x")
    ax.grid(axis="y", visible=False)
    fig.savefig(os.path.join(CH_DIR, "C2_producers_share.png"))
    plt.close(fig)

    # C3 - payment mix --------------------------------------------------------
    df = tmap["T5_Payment_Methods"]
    df = df[df["Actor"].isin(["Aratdar", "Bepari/Faria", "Retailer"])]
    fig, ax = plt.subplots(figsize=(8.5, 5.5), constrained_layout=True)
    xp = np.arange(len(df))
    bottom = np.zeros(len(df))
    for lab, col, colr in (("Cash", "Cash_pct_mean", "#8c8c8c"),
                           ("MFS (bKash/Nagad)", "MFS_pct_mean", "#2e8b57"),
                           ("Credit", "Credit_pct_mean", "#c0504d")):
        vals = df[col].fillna(0).to_numpy()
        ax.bar(xp, vals, 0.55, bottom=bottom, label=lab, color=colr,
               edgecolor="white", linewidth=0.4)
        for i, (v, b) in enumerate(zip(vals, bottom)):
            if v >= 8:
                ax.text(i, b + v / 2, f"{v:.0f}%", ha="center", va="center",
                        color="white", fontsize=9)
        bottom += vals
    ax.set_xticks(xp)
    ax.set_xticklabels(df["Actor"])
    ax.set_ylabel("Mean share of receipts (%)")
    ax.set_ylim(0, 100)
    ax.set_title("How channel actors get paid (mean share of receipts)")
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), ncol=3)
    ax.grid(axis="x", visible=False)
    fig.savefig(os.path.join(CH_DIR, "C3_payment_mix.png"))
    plt.close(fig)

    # C4 - education -----------------------------------------------------------
    t1 = tmap["T1_Respondent_Profile"]
    t1 = t1[t1["Actor"] != "All traders"]
    edu_cols = [c for c in t1.columns if c.startswith("Edu_")]
    fig, ax = plt.subplots(figsize=(9.5, 5.5), constrained_layout=True)
    xp = np.arange(len(edu_cols))
    w = 0.2
    for i, (_, r) in enumerate(t1.iterrows()):
        ax.bar(xp + (i - 1.5) * w, [r[c] or 0 for c in edu_cols], w,
               label=r["Actor"], color=[C_PRODUCER, C_ARATDAR, C_BEPARI, C_RETAIL][i],
               edgecolor="white", linewidth=0.4)
    ax.set_xticks(xp)
    ax.set_xticklabels([c.replace("Edu_", "").replace("_pct", "").replace("_", " ")
                        for c in edu_cols])
    ax.set_ylabel("% of respondents")
    ax.set_title("Education level of trading actors")
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), ncol=4)
    ax.grid(axis="y")
    ax.grid(axis="x", visible=False)
    fig.savefig(os.path.join(CH_DIR, "C4_education.png"))
    plt.close(fig)

    # C5 - problems -------------------------------------------------------------
    df = tmap["T6_Problems"].head(12).sort_values("Total_pct")
    fig, ax = plt.subplots(figsize=(10, 6.5), constrained_layout=True)
    ax.barh(df["Problem"], df["Total_pct"], color=C_ARATDAR, alpha=0.9)
    for i, (v, n_) in enumerate(zip(df["Total_pct"], df["Total_n"])):
        ax.text(v + 0.6, i, f"{v:.0f}% (n={n_})", va="center", fontsize=8.5)
    ax.set_xlabel("% of traders mentioning (multi-response, n=90)")
    ax.set_xlim(0, max(df["Total_pct"]) * 1.22)
    ax.set_title("Leading marketing problems reported by traders")
    ax.grid(axis="x")
    ax.grid(axis="y", visible=False)
    fig.savefig(os.path.join(CH_DIR, "C5_problems.png"))
    plt.close(fig)

    # C6 - retail price heatmap --------------------------------------------------
    df = tmap["T11_Retail_Price_by_Market"].set_index("Species")
    mkt_names = list(df.columns)
    mat = df.to_numpy(dtype=float)
    fig, ax = plt.subplots(figsize=(9.5, 6.5), constrained_layout=True)
    im = ax.imshow(np.ma.masked_invalid(mat), cmap="YlGnBu", aspect="auto")
    ax.set_xticks(range(len(mkt_names)))
    ax.set_xticklabels(mkt_names, rotation=25, ha="right")
    ax.set_yticks(range(len(df.index)))
    ax.set_yticklabels(df.index)
    for i in range(mat.shape[0]):
        for j in range(mat.shape[1]):
            if not np.isnan(mat[i, j]):
                ax.text(j, i, f"{mat[i, j]:.0f}", ha="center", va="center",
                        fontsize=8.5,
                        color="white" if mat[i, j] > np.nanmax(mat) * 0.55 else "#263238")
    cb = fig.colorbar(im, ax=ax, shrink=0.85)
    cb.set_label("BDT per kg")
    ax.set_title("Retailer selling prices by species and market (BDT/kg)")
    ax.grid(visible=False)
    fig.savefig(os.path.join(CH_DIR, "C6_retail_price_heatmap.png"))
    plt.close(fig)

    # C7 - daily capacity boxplot -------------------------------------------------
    data_by, labels = [], []
    for label, rs, key in (("Aratdar", d["A"], "Daily_capacity_kg"),
                           ("Bepari/Faria", d["B"], "Daily_capacity_kg"),
                           ("Retailer", d["R"], "Daily_capacity_kg")):
        vals = [num(r.get(key)) for r in rs]
        data_by.append([v for v in vals if v is not None])
        labels.append(f"{label}\n(n={len(data_by[-1])})")
    fig, ax = plt.subplots(figsize=(8.5, 5.5), constrained_layout=True)
    bp = ax.boxplot(data_by, tick_labels=labels, patch_artist=True, widths=0.5,
                    medianprops=dict(color="#111111", linewidth=1.5))
    for patch, colr in zip(bp["boxes"], [C_PRODUCER, C_ARATDAR, C_RETAIL]):
        patch.set_facecolor(colr)
        patch.set_alpha(0.75)
    ax.set_yscale("log")
    ax.set_ylabel("Daily volume handled (kg, log scale)")
    ax.set_title("Scale gap between channel actors (daily volume)")
    ax.grid(axis="y")
    ax.grid(axis="x", visible=False)
    fig.savefig(os.path.join(CH_DIR, "C7_daily_capacity.png"))
    plt.close(fig)


# ----------------------------------------------------------------------------
# Analysis_Summary.xlsx writer
# ----------------------------------------------------------------------------
def write_summary_xlsx(tables, path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    idx = wb.create_sheet("Index")
    idx.cell(row=2, column=2, value="MS-499 Analysis Summary - Marine Fish "
              "Marketing, Chattogram 2026").font = Font(bold=True, size=14,
                                                        color=HDR_FILL)
    idx.cell(row=3, column=2, value="Every table also saved as CSV in "
              "analysis_outputs/tables/; figures in analysis_outputs/charts/"
              ).font = Font(italic=True, size=9, color="808080")
    idx_hdr = ["Sheet", "Table title", "Rows", "Cols"]
    for j, h in enumerate(idx_hdr):
        c = idx.cell(row=4, column=2 + j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HDR_FILL)
    idx.column_dimensions["B"].width = 26
    idx.column_dimensions["C"].width = 52
    idx.column_dimensions["D"].width = 8
    idx.column_dimensions["E"].width = 8
    idx.sheet_view.showGridLines = False

    for t_i, (sheet, title, df, note) in enumerate(tables):
        ws = wb.create_sheet(sheet[:31])
        ws.cell(row=2, column=2, value=title).font = Font(bold=True, size=13,
                                                           color=HDR_FILL)
        ws.cell(row=3, column=2, value=note).font = Font(italic=True, size=9,
                                                          color="808080")
        ws.merge_cells(start_row=3, start_column=2, end_row=3,
                       end_column=min(len(df.columns) + 1, 10))
        headers = list(df.columns)
        for j, h in enumerate(headers):
            c = ws.cell(row=4, column=2 + j, value=str(h))
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor=HDR_FILL)
            c.alignment = Alignment(wrap_text=True, vertical="center",
                                    horizontal="center")
        for i, (_, row) in enumerate(df.iterrows()):
            for j, h in enumerate(headers):
                v = row[h]
                if pd.isna(v):
                    v = None
                elif isinstance(v, (np.integer,)):
                    v = int(v)
                elif isinstance(v, (np.floating,)):
                    v = float(v)
                c = ws.cell(row=5 + i, column=2 + j, value=v)
                if isinstance(v, float):
                    c.number_format = "0.0" if abs(v) < 10000 else "#,##0"
                if i % 2 == 1:
                    c.fill = PatternFill("solid", fgColor=BAND_FILL)
                if isinstance(v, str) and len(v) > 30:
                    c.alignment = Alignment(horizontal="left")
        ws.freeze_panes = "C5"
        for j, h in enumerate(headers):
            w_col = max([len(str(h)) * 1.05] +
                        [len(str(row[h])) for _, row in df.iterrows()
                         if pd.notna(row[h])])
            ws.column_dimensions[get_column_letter(2 + j)].width = max(9,
                                                                       min(w_col + 2, 40))
        ws.sheet_view.showGridLines = False
        for j, v in enumerate((sheet, title, len(df), len(df.columns))):
            idx.cell(row=5 + t_i, column=2 + j, value=v)
    wb.properties.creator = "Z.ai"
    wb.save(path)


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="MS-499 survey analysis pipeline")
    ap.add_argument("--data", default=None, help="path to filled workbook (.xlsx)")
    args = ap.parse_args()
    path = args.data or find_default_data()
    print(f"Reading workbook: {path}")
    d = load_data(path)
    print(f"Loaded: A={len(d['A'])}, B={len(d['B'])}, R={len(d['R'])}, "
          f"C={len(d['C'])}, price_obs={len(d['PO'])}, purchases={len(d['CPF'])}")
    print(f"Survey window: {d['window'][0]} - {d['window'][1]}")

    os.makedirs(TBL_DIR, exist_ok=True)
    os.makedirs(CH_DIR, exist_ok=True)

    tables = build_tables(d)
    for sheet, title, df, note in tables:
        csv = os.path.join(TBL_DIR, f"{sheet}.csv")
        df.to_csv(csv, index=False, encoding="utf-8-sig")
        print(f"  table  {sheet}.csv  ({len(df)} rows)")

    write_summary_xlsx(tables, XLSX_OUT)
    print(f"  workbook {XLSX_OUT}")

    make_charts(d, tables)
    own = {f"C{i}_" for i in range(1, 8)}
    for f in sorted(os.listdir(CH_DIR)):
        if f.startswith(tuple(own)):
            print(f"  chart  {f}")
    print("DONE - outputs in analysis_outputs/")


if __name__ == "__main__":
    main()
