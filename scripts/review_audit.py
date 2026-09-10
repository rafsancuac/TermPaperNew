#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MS-499 Term Paper - SUPERVISOR-LEVEL DATA AUDIT (beyond template QC)

Checks logical consistency, plausibility, chain coherence and internal
agreement of the filled workbook, and independently recomputes the
headline statistics to validate the analysis pipeline.

Output: console report with [PASS]/[WARN]/[FAIL] per check.
Usage:  python scripts/review_audit.py [--data path/to/filled.xlsx]
"""
import argparse
import os
from collections import Counter, defaultdict
from datetime import date, datetime

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DEFAULT = os.path.join(ROOT, "04_data_filled",
                       "Marine_Fish_Marketing_Data_Entry_Chattogram_Filled.xlsx")
MAUND = 37.32

MARKETS = ["M1", "M2", "M3", "M4", "M5", "M6"]
SPECIES = {
    "S01": "Ilish", "S02": "Rupchanda", "S03": "Lakkha", "S04": "Koral",
    "S05": "Surma", "S06": "Churi", "S07": "Poa", "S08": "Kankoita",
    "S09": "Loitta", "S10": "Harina",
}
# Realistic 2026 Chattogram RETAIL bands, BDT/kg (loose, for realism flags)
RETAIL_BAND = {
    "S01": (900, 2200), "S02": (800, 1900), "S03": (600, 1400),
    "S04": (550, 1300), "S05": (300, 750), "S06": (220, 650),
    "S07": (280, 700), "S08": (180, 550), "S09": (120, 420),
    "S10": (100, 380),
}
# survey schedule: market -> allowed interview dates
SCHEDULE = {
    "M1": {date(2026, 3, 2), date(2026, 3, 3)},
    "M2": {date(2026, 3, 4)}, "M3": {date(2026, 3, 5)},
    "M4": {date(2026, 3, 6)}, "M5": {date(2026, 3, 7)},
    "M6": {date(2026, 3, 3)},
}

n_pass = n_warn = n_fail = 0


def rep(status, name, detail=""):
    global n_pass, n_warn, n_fail
    tag = {"PASS": "[PASS]", "WARN": "[WARN]", "FAIL": "[FAIL]"}[status]
    if status == "PASS":
        n_pass += 1
    elif status == "WARN":
        n_warn += 1
    else:
        n_fail += 1
    print(f"{tag} {name}" + (f"  --  {detail}" if detail else ""))


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def dstr(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def read_sheet(ws, id_col=2, require_any=None):
    headers = {}
    for c in range(2, ws.max_column + 1):
        h = ws.cell(row=4, column=c).value
        if h is not None and str(h).strip():
            headers[c] = str(h).strip()
    rows = []
    for r in range(5, ws.max_row + 1):
        rid = ws.cell(row=r, column=id_col).value
        if rid in (None, ""):
            continue
        rec = {h: ws.cell(row=r, column=c).value for c, h in headers.items()}
        # skip unused spare rows (prefilled ID but no entered data)
        if require_any and all(rec.get(k) in (None, "") for k in require_any):
            continue
        rows.append(rec)
    return rows


def price_kg(raw, unit, cached):
    v = num(cached)
    if v is not None:
        return v
    v = num(raw)
    if v is None:
        return None
    return v / MAUND if str(unit or "").strip() == "Maund" else v


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    A = read_sheet(wb["Form_A_Aratdar"], require_any=["Interview_Date", "Age_years"])
    B = read_sheet(wb["Form_B_Bepari_Faria"], require_any=["Interview_Date", "Age_years"])
    R = read_sheet(wb["Form_R_Khuchra"], require_any=["Interview_Date", "Age_years"])
    C = read_sheet(wb["Form_C_Consumer"], require_any=["Interview_Date", "Visit_frequency"])
    PO = read_sheet(wb["Price_Observations"])
    CPF = read_sheet(wb["Consumer_Purchases_Focal"])
    COF = read_sheet(wb["Consumer_Other_Fish"])
    FM = read_sheet(wb["Form_M_Market_Observation"])
    LOG = read_sheet(wb["Data_Collection_Log"])
    TAG = read_sheet(wb["Tag_Price_Sheet"])

    print("=" * 78)
    print("SUPERVISOR-LEVEL AUDIT —", os.path.basename(path))
    print("=" * 78)

    # ---------------- 1. quota & ID structure ----------------
    for name, rows, suf in [("Aratdar", A, "A"), ("Bepari/Faria", B, "B"),
                            ("Retail", R, "R"), ("Consumer", C, "C")]:
        cnt = Counter(r["Market"] for r in rows)
        bad = [m for m in MARKETS if cnt.get(m, 0) != 5]
        rep("PASS" if not bad else "FAIL", f"Quota 5x6 {name}",
            f"per-market={dict(sorted(cnt.items()))}" if bad else "30 rows, 5 per market")
        ids = [r["Respondent_ID"] for r in rows]
        pat_bad = [i for i in ids if not i.startswith(tuple(f"{m}-{suf}" for m in MARKETS))]
        rep("PASS" if not pat_bad else "FAIL", f"ID pattern {name}",
            f"bad={pat_bad[:3]}" if pat_bad else "all match M#-X-##")
        dup = [i for i, k in Counter(ids).items() if k > 1]
        rep("PASS" if not dup else "FAIL", f"No duplicate IDs {name}",
            str(dup) if dup else "")
    rep("PASS" if len(A) + len(B) + len(R) + len(C) == 120 else "FAIL",
        "Total respondents = 120", f"got {len(A)+len(B)+len(R)+len(C)}")

    # ---------------- 2. interview dates vs schedule ----------------
    allrows = [("A", A), ("B", B), ("R", R), ("C", C)]
    mism = []
    for _, rows in allrows:
        for r in rows:
            d = dstr(r.get("Interview_Date"))
            m = r.get("Market")
            if d is None or d not in SCHEDULE.get(m, set()):
                mism.append((r.get("Respondent_ID"), m, d))
    rep("PASS" if not mism else "FAIL", "Interview dates match field schedule",
        f"mismatches={mism[:5]}" if mism else "all 120 within scheduled market-days")

    log_days = [dstr(r.get("Date")) for r in LOG]
    wk_ok = all(
        (dstr(r.get("Date")).weekday() == ["mon", "tue", "wed",
         "thu", "fri", "sat", "sun"].index(str(r.get("Weekday", "")).strip()[:3].lower()))
        for r in LOG if dstr(r.get("Date")) and r.get("Weekday"))
    rep("PASS" if wk_ok else "FAIL", "Log weekday vs date consistency",
        "" if wk_ok else "weekday labels do not match calendar dates")
    rep("PASS" if len(LOG) == 7 else "WARN", "Log has 7 visits (pilot + 6 days)",
        f"got {len(LOG)}")

    # ---------------- 3. demographic plausibility ----------------
    for name, rows in [("Aratdar", A), ("Bepari", B), ("Retail", R)]:
        ages = [num(r["Age_years"]) for r in rows if num(r["Age_years"])]
        exps = [num(r["Years_in_business"]) for r in rows if num(r["Years_in_business"])]
        out_age = [(r["Respondent_ID"], r["Age_years"]) for r in rows
                   if not (18 <= (num(r["Age_years"]) or 0) <= 80)]
        rep("PASS" if not out_age else "FAIL", f"Age 18-80 {name}",
            str(out_age[:3]) if out_age else f"mean={np.mean(ages):.1f}")
        bad_exp = [(r["Respondent_ID"], r["Age_years"], r["Years_in_business"])
                   for r in rows
                   if num(r["Age_years"]) and num(r["Years_in_business"])
                   and num(r["Years_in_business"]) > num(r["Age_years"]) - 14]
        rep("PASS" if not bad_exp else "FAIL", f"Experience <= Age-14 {name}",
            str(bad_exp[:3]) if bad_exp else "all coherent")

    edu_age_bad = []
    for _, rows in allrows:
        for r in rows:
            a = num(r.get("Age_years"))
            e = r.get("Education")
            if a and e == "Higher_Secondary" and a < 19:
                edu_age_bad.append((r["Respondent_ID"], a, e))
            if a and e == "Secondary" and a < 17:
                edu_age_bad.append((r["Respondent_ID"], a, e))
    rep("PASS" if not edu_age_bad else "WARN", "Education level vs age",
        str(edu_age_bad[:5]) if edu_age_bad else "all plausible")
    fam_bad = [(r["Respondent_ID"], r["Family_members"]) for _, rows in
               [("A", A), ("B", B), ("R", R)] for r in rows
               if not (1 <= (num(r.get("Family_members")) or 5) <= 15)]
    rep("PASS" if not fam_bad else "WARN", "Family size 1-15",
        str(fam_bad[:3]) if fam_bad else "plausible")

    # ---------------- 4. business-scale ordering ----------------
    cap = {k: np.mean([num(r["Daily_capacity_kg"]) for r in rows
                       if num(r.get("Daily_capacity_kg"))])
           for k, rows in [("Aratdar", A), ("Bepari", B), ("Retail", R)]}
    ok = cap["Aratdar"] > cap["Bepari"] > cap["Retail"]
    rep("PASS" if ok else "FAIL", "Volume ordering A > B > R",
        f"aratdar={cap['Aratdar']:.0f} > bepari={cap['Bepari']:.0f} "
        f"> retail={cap['Retail']:.0f} kg/day")
    bsub = defaultdict(list)
    for r in B:
        v = num(r.get("Daily_capacity_kg"))
        if v:
            bsub[r.get("Actor_subtype")].append(v)
    if "Bepari" in bsub and "Faria" in bsub:
        ok2 = np.mean(bsub["Bepari"]) > np.mean(bsub["Faria"])
        rep("PASS" if ok2 else "WARN", "Bepari capacity > Faria",
            f"bepari={np.mean(bsub['Bepari']):.0f}, faria={np.mean(bsub['Faria']):.0f}")

    # ---------------- 5. payment shares ----------------
    pay_bad = []
    for name, rows in [("A", A), ("B", B), ("R", R)]:
        for r in rows:
            s = sum(num(r.get(k)) or 0 for k in
                    ("Payment_cash_pct", "Payment_MFS_pct", "Payment_credit_pct"))
            if abs(s - 100) > 0.51:
                pay_bad.append((name, r["Respondent_ID"], s))
    rep("PASS" if not pay_bad else "FAIL", "Payment pct sums to 100 (A/B/R)",
        str(pay_bad[:5]) if pay_bad else "all 90 respondents")

    mfs_bad = []
    for name, rows in [("A", A), ("B", B), ("R", R)]:
        for r in rows:
            if (num(r.get("Payment_MFS_pct")) or 0) == 0 and \
                    r.get("MFS_transaction_frequency") == "Regular":
                mfs_bad.append((name, r["Respondent_ID"]))
    rep("PASS" if not mfs_bad else "WARN", "MFS% > 0 iff frequency used",
        str(mfs_bad[:5]) if mfs_bad else "consistent")

    # ---------------- 6. price chain coherence ----------------
    # Convention (since supervisor-review commit): producer & auction sell
    # observed only at landing markets M1/M6 (first-sale proxy, Methodology
    # 3.11c); consumer price = consumer-paid (Form C focal purchases);
    # retail margin anchored to consumer-paid price.
    LANDING = {"M1", "M6"}
    chain = defaultdict(lambda: defaultdict(list))  # species -> stage -> [prices]
    for p in PO:
        sp = p.get("Species_code")
        at = p.get("Actor_type")
        mk = str(p.get("Market")).strip()
        if sp is None:
            continue
        bp = price_kg(p.get("Buy_price_raw"), p.get("Unit"), p.get("Buy_BDT_per_kg"))
        sp_ = price_kg(p.get("Sell_price_raw"), p.get("Unit"), p.get("Sell_BDT_per_kg"))
        if at == "Aratdar" and bp and mk in LANDING:
            chain[sp]["producer"].append(bp)
            if sp_:
                chain[sp]["aratdar_sell"].append(sp_)
        if at == "Bepari_Faria" and bp:
            chain[sp]["bepari_buy"].append(bp)
            if sp_:
                chain[sp]["bepari_sell"].append(sp_)
        if at == "Khuchra" and bp:
            chain[sp]["retail_buy"].append(bp)
            if sp_:
                chain[sp]["retail_sell"].append(sp_)
    for r in CPF:
        if r.get("Purchased_today") == "Yes" and r.get("Species_code"):
            v = price_kg(r.get("Price_raw"), r.get("Unit"), r.get("Price_BDT_per_kg"))
            if v:
                chain[r["Species_code"]]["consumer_paid"].append(v)

    mono_bad, share = [], {}
    for sp in SPECIES:
        c = chain[sp]
        prod = np.mean(c["producer"]) if c["producer"] else None
        asel = np.mean(c["aratdar_sell"]) if c["aratdar_sell"] else None
        bsel = np.mean(c["bepari_sell"]) if c["bepari_sell"] else None
        cons = np.mean(c["consumer_paid"]) if c["consumer_paid"] else None
        rsell = np.mean(c["retail_sell"]) if c["retail_sell"] else None
        if prod and cons:
            share[sp] = 100 * prod / cons
        vals = [("producer", prod), ("aratdar_sell", asel),
                ("bepari_sell", bsel), ("consumer_paid", cons)]
        for (n1, v1), (n2, v2) in zip(vals, vals[1:]):
            if v1 and v2 and v2 <= v1:
                mono_bad.append((sp, n1, round(v1), n2, round(v2)))
    rep("PASS" if not mono_bad else "FAIL",
        "Chain monotonic: producer < aratdar < bepari < consumer (per species)",
        str(mono_bad[:4]) if mono_bad else "increasing wherever both stages exist")
    if share:
        overall = np.mean(list(share.values()))
        rep("PASS" if 55 < overall < 85 else "WARN",
            "Producer's share recomputed (consumer-anchored, all species)",
            f"species-mean = {overall:.1f}% "
            f"(range {min(share.values()):.0f}-{max(share.values()):.0f}%)")

    # buy<=sell within each PO row
    row_bad = [p.get("Obs_ID") for p in PO
               if price_kg(p.get("Buy_price_raw"), p.get("Unit"), p.get("Buy_BDT_per_kg"))
               and price_kg(p.get("Sell_price_raw"), p.get("Unit"), p.get("Sell_BDT_per_kg"))
               and price_kg(p.get("Sell_price_raw"), p.get("Unit"), p.get("Sell_BDT_per_kg"))
               < price_kg(p.get("Buy_price_raw"), p.get("Unit"), p.get("Buy_BDT_per_kg"))]
    rep("PASS" if not row_bad else "FAIL", "Sell >= Buy within each observation",
        str(row_bad[:5]) if row_bad else f"all {len(PO)} rows")

    # margin sanity per stage (new convention; chain-complete species only -
    # completeness matches the pipeline rule: all four stages present AND at
    # least 3 consumer-paid observations, MIN_CONS in run_analysis.py)
    mgn = defaultdict(list)
    chain_complete = []
    for sp in SPECIES:
        c = chain[sp]
        if not (c["producer"] and c["aratdar_sell"] and c["bepari_sell"]
                and c["consumer_paid"]) or len(c["consumer_paid"]) < 3:
            continue
        chain_complete.append(sp)
        mgn["aratdar"].append(np.mean(c["aratdar_sell"]) - np.mean(c["producer"]))
        mgn["bepari"].append(np.mean(c["bepari_sell"]) - np.mean(c["aratdar_sell"]))
        mgn["retail"].append(np.mean(c["consumer_paid"]) - np.mean(c["bepari_sell"]))
    for st in ("aratdar", "bepari", "retail"):
        v = np.mean(mgn[st]) if mgn[st] else 0
        rep("PASS" if 0 < v < 200 else "WARN", f"Margin {st} plausible (BDT/kg)",
            f"species-mean = {v:.1f} (chain-complete: {len(chain_complete)} species)")
    if chain_complete:
        sum_m = np.mean(mgn["aratdar"]) + np.mean(mgn["bepari"]) + np.mean(mgn["retail"])
        spread = np.mean([np.mean(chain[sp]["consumer_paid"])
                          - np.mean(chain[sp]["producer"]) for sp in chain_complete])
        rep("PASS" if abs(sum_m - spread) < 0.6 else "FAIL",
            "Margins telescope: A+B+R = spread (100% accounting)",
            f"A+B+R = {sum_m:.2f} vs spread = {spread:.2f} BDT/kg")

    # ---------------- 7. retail price realism bands ----------------
    out_band = []
    for p in PO:
        sp = p.get("Species_code")
        if sp not in RETAIL_BAND or p.get("Actor_type") != "Khuchra":
            continue
        pr = price_kg(p.get("Sell_price_raw"), p.get("Unit"), p.get("Sell_BDT_per_kg"))
        if pr and not (RETAIL_BAND[sp][0] * 0.75 <= pr <= RETAIL_BAND[sp][1] * 1.35):
            out_band.append((p.get("Obs_ID"), sp, round(pr)))
    rep("PASS" if not out_band else "WARN", "Retail prices inside realistic 2026 band",
        f"outliers={out_band[:5]} ({len(out_band)} rows)" if out_band else
        "all khuchra sell-prices within bands")

    # consumer-paid vs retailer-sell gap (independent collection)
    gap_bad = []
    for sp in SPECIES:
        con = [price_kg(r.get("Price_raw"), r.get("Unit"), r.get("Price_BDT_per_kg"))
               for r in CPF if r.get("Species_code") == sp
               and r.get("Purchased_today") == "Yes"]
        con = [x for x in con if x]
        rs = chain[sp]["retail_sell"]
        if con and rs:
            g = (np.mean(con) - np.mean(rs)) / np.mean(rs) * 100
            if abs(g) > 10:
                gap_bad.append((sp, f"{g:+.1f}%"))
    rep("PASS" if not gap_bad else "WARN",
        "Consumer-paid vs retailer-quoted prices (independent), |gap|<=10%",
        str(gap_bad) if gap_bad else "all species within 10%")

    # ---------------- 8. unit usage ----------------
    u_wh = Counter(str(p.get("Unit")).strip() for p in PO
                   if p.get("Actor_type") in ("Aratdar", "Bepari_Faria"))
    u_rt = Counter(str(p.get("Unit")).strip() for p in PO
                   if p.get("Actor_type") == "Khuchra")
    wh_ok = u_wh.get("Maund", 0) >= u_wh.get("Kg", 0)
    rt_ok = u_rt.get("Kg", 0) >= u_rt.get("Maund", 0)
    rep("PASS" if wh_ok else "WARN", "Wholesale quotes mostly Maund",
        str({k: v for k, v in u_wh.items() if k != 'None'}))
    rep("PASS" if rt_ok else "WARN", "Retail quotes mostly Kg",
        str({k: v for k, v in u_rt.items() if k != 'None'}))

    # ---------------- 9. Pair_ID transaction consistency ----------------
    pairs = defaultdict(list)
    for p in PO:
        pid = p.get("Pair_ID")
        if pid and str(pid).strip().upper() not in ("", "K", "D", "NONE", "-"):
            pairs[str(pid).strip()].append(p)
    pair_bad, diffs = [], []
    npairs = 0
    # channel order: buyer must sit at the same or a downstream stage of the
    # seller (Aratdar -> Bepari_Faria -> Khuchra). The seller's own row also
    # carries a buy price, so buyer rows must be a DIFFERENT respondent.
    ACTOR_ORDER = {"Aratdar": 0, "Bepari_Faria": 1, "Khuchra": 2}
    for pid, rows in pairs.items():
        if len(rows) < 2:
            continue
        npairs += 1
        best = None  # (gap, s_price, b_price)
        for a in range(len(rows)):
            for b2 in range(len(rows)):
                if a == b2:
                    continue
                s_r, b_r = rows[a], rows[b2]
                s = price_kg(s_r.get("Sell_price_raw"), s_r.get("Unit"),
                             s_r.get("Sell_BDT_per_kg"))
                bb = price_kg(b_r.get("Buy_price_raw"), b_r.get("Unit"),
                              b_r.get("Buy_BDT_per_kg"))
                if s is None or bb is None:
                    continue
                if ACTOR_ORDER.get(s_r.get("Actor_type"), 9) > \
                        ACTOR_ORDER.get(b_r.get("Actor_type"), 9):
                    continue
                gap = abs(s - bb) / max(s, bb)
                if best is None or gap < best[0]:
                    best = (gap, s, bb)
        if best is None:
            continue
        diffs.append(best[0])
        if best[0] > 0.05:
            pair_bad.append((pid, round(best[1]), round(best[2])))
    rep("PASS" if len(diffs) >= 10 else "WARN",
        f"Matched pairs available for Wilcoxon (n={npairs})",
        f"usable price-pairs={len(diffs)}")
    if diffs:
        med = np.median(diffs) * 100
        rep("PASS" if med < 2 else "WARN",
            "Pair prices: seller-sell ~ buyer-buy (same transaction)",
            f"median rel. gap {med:.2f}%, worst={pair_bad[:3] if pair_bad else 'none'}")
    pm = Counter()
    for _, rows in allrows:
        for r in rows:
            pid = str(r.get("Pair_ID") or "").strip()
            if pid and pid.upper() not in ("", "K", "D", "NONE", "-"):
                parts = pid.split("-")
                pm[parts[1] if len(parts) > 1 else "?"] += 1
    rep("PASS" if all(pm.get(m, 0) >= 2 for m in MARKETS) else "WARN",
        "Pairs spread across markets", str(dict(sorted(pm.items()))))

    # ---------------- 10. consumer sheets ----------------
    cp_bad = [r.get("Respondent_ID") for r in CPF
              if r.get("Purchased_today") == "No"
              and num(r.get("Price_raw")) is not None]
    rep("PASS" if not cp_bad else "FAIL", "Focal purchases: No -> no price",
        str(cp_bad[:5]) if cp_bad else "consistent")
    kn = sum(1 for r in CPF if r.get("Purchased_today") == "No")
    rep("PASS", "Focal purchases Yes/No mix", f"No={kn}, Yes={len(CPF)-kn}")
    con_prices = [price_kg(r.get("Price_raw"), r.get("Unit"), r.get("Price_BDT_per_kg"))
                  for r in CPF if r.get("Purchased_today") == "Yes"]
    if con_prices:
        lo = [x for x in con_prices if x and x < 60]
        rep("PASS" if not lo else "WARN", "Consumer prices sensible (>=60 BDT/kg)",
            f"suspect={lo[:5]}" if lo else f"mean={np.mean(con_prices):.0f}")
    bf = Counter(r.get("Bought_from") for r in CPF if r.get("Bought_from"))
    rep("PASS", "Bought_from categories", str(dict(bf)))
    vm = Counter(r.get("Visit_frequency") for r in C)
    rep("PASS", "Consumer visit frequency", str(dict(vm)))
    cons_pay = Counter(r.get("Payment_method") for r in C)
    rep("PASS", "Consumer payment methods", str(dict(cons_pay)))

    # consumer focal ilish vs retail ilish
    il_con = [price_kg(r.get("Price_raw"), r.get("Unit"), r.get("Price_BDT_per_kg"))
              for r in CPF if r.get("Species_code") == "S01"
              and r.get("Purchased_today") == "Yes"]
    il_rt = chain["S01"]["retail_sell"]
    if il_con and il_rt:
        ratio = np.mean(il_con) / np.mean(il_rt)
        rep("PASS" if 0.85 <= ratio <= 1.25 else "WARN",
            "Consumer-paid Ilish ~ retail sell price (S01)",
            f"consumer {np.mean(il_con):.0f} vs retail {np.mean(il_rt):.0f} "
            f"(ratio {ratio:.2f})")

    # other fish: names not duplicating focal species
    focal_names = set(v.lower() for v in SPECIES.values())
    dup_names = [r.get("Fish_name_local") for r in COF
                 if str(r.get("Fish_name_local", "")).strip().lower() in focal_names]
    rep("PASS" if not dup_names else "WARN",
        "Other-fish sheet: no focal-species duplicates",
        str(dup_names[:5]) if dup_names else
        f"{len(COF)} rows, e.g. {sorted(set(str(r['Fish_name_local']) for r in COF))[:8]}")
    ofp = [price_kg(r.get("Price_raw"), r.get("Unit"), r.get("Price_BDT_per_kg"))
           for r in COF]
    if ofp:
        # species-aware bands: shrimp / lobster / crab legitimately trade
        # above the generic 60-600 BDT/kg band of common table fish
        HIGH_BAND = {"Lobster": 3600, "Bagda": 1400, "Kakra": 750}
        def _of_cap(r):
            name = str(r.get("Fish_name_local") or "")
            for k, cap in HIGH_BAND.items():
                if k.lower() in name.lower():
                    return cap
            return 600
        bad_of = [x for x, r in zip(ofp, COF)
                  if x and (x < 60 or x > _of_cap(r))]
        rep("PASS" if not bad_of else "WARN", "Other-fish prices in 60-600 band (+shrimp/lobster)",
            f"suspect={bad_of[:5]}" if bad_of else f"mean={np.mean(ofp):.0f} BDT/kg")

    # ---------------- 11. Form M ----------------
    FM = [r for r in FM if str(r.get("Market")).strip() in MARKETS]
    fm_markets = sorted(str(r.get("Market")).strip() for r in FM)
    rep("PASS" if fm_markets == sorted(MARKETS) else "FAIL",
        "Form M: one row per market", str(fm_markets))
    gps_bad = []
    for r in FM:
        g = str(r.get("GPS_coordinates", ""))
        try:
            lat = float(g.split(",")[0].strip())
            lon = float(g.split(",")[1].strip())
            if not (22.1 <= lat <= 22.5 and 91.6 <= lon <= 92.0):
                gps_bad.append((r.get("Market"), g))
        except Exception:
            gps_bad.append((r.get("Market"), g))
    rep("PASS" if not gps_bad else "WARN", "GPS within Chattogram bounds",
        str(gps_bad) if gps_bad else "all 6 in 22.1-22.5N, 91.6-92.0E")
    fm_date_bad = []
    for r in FM:
        d = dstr(r.get("Obs_Date"))
        if d not in SCHEDULE.get(str(r.get("Market")).strip(), set()):
            fm_date_bad.append((r.get("Market"), d))
    rep("PASS" if not fm_date_bad else "WARN", "Form M dates match schedule",
        str(fm_date_bad) if fm_date_bad else "all 6 on scheduled days")

    # ---------------- 12. tag-price vs retail observations ----------------
    tag_by_msp = defaultdict(list)
    for t in TAG:
        key = (str(t.get("Market")).strip(), str(t.get("Fish_name_local")).strip().lower())
        v = price_kg(t.get("Price_raw"), t.get("Unit"), t.get("Price_BDT_per_kg"))
        if v:
            tag_by_msp[key].append(v)
    # map local fish names to species codes
    name2code = {v.lower(): k for k, v in SPECIES.items()}
    retail_ms = defaultdict(list)  # (market, species) -> khuchra sell prices
    for p in PO:
        mk, code = str(p.get("Market")).strip(), p.get("Species_code")
        if p.get("Actor_type") == "Khuchra" and code in SPECIES:
            rp = price_kg(p.get("Sell_price_raw"), p.get("Unit"),
                          p.get("Sell_BDT_per_kg"))
            if rp:
                retail_ms[(mk, code)].append(rp)
    mism_cnt, checked, gaps = 0, 0, []
    for (mk, name), vals in tag_by_msp.items():
        code = name2code.get(name)
        if not code or (mk, code) not in retail_ms:
            continue
        checked += 1
        gap = (np.mean(vals) - np.mean(retail_ms[(mk, code)])) / np.mean(retail_ms[(mk, code)])
        gaps.append(abs(gap) * 100)
        if abs(gap) > 0.20:
            mism_cnt += 1
    rep("PASS" if checked >= 10 and mism_cnt == 0 else "WARN",
        "Tag-price vs khuchra observations (same market+species)",
        f"{mism_cnt}/{checked} pairs deviate >20%; "
        f"median gap {np.median(gaps):.1f}%" if gaps else "no overlap found")
    tag_dates = [dstr(t.get("Date")) for t in TAG if dstr(t.get("Date"))]
    win = {date(2026, 3, d) for d in range(2, 8)}
    rep("PASS" if all(d in win for d in tag_dates) else "WARN",
        "Tag-price dates within survey window",
        f"{min(tag_dates)}..{max(tag_dates)}" if tag_dates else "none")

    # ---------------- 13. species coverage ----------------
    cov = Counter(p.get("Species_code") for p in PO if p.get("Species_code"))
    low = [f"{k}:{v}" for k, v in sorted(cov.items()) if v < 20]
    rep("PASS" if not low else "WARN", "Species observation counts (>=20 each)",
        f"low={low}" if low else str(dict(sorted(cov.items()))))
    at_cov = Counter((p.get("Actor_type"), p.get("Species_code")) for p in PO
                     if price_kg(p.get("Buy_price_raw"), p.get("Unit"),
                                 p.get("Buy_BDT_per_kg")) or
                     price_kg(p.get("Sell_price_raw"), p.get("Unit"),
                              p.get("Sell_BDT_per_kg")))
    missing_stage = [(at, sp) for at in ("Aratdar", "Bepari_Faria", "Khuchra")
                     for sp in SPECIES if at_cov.get((at, sp), 0) == 0]
    rep("PASS" if not missing_stage else "WARN",
        "Every species observed at every chain stage",
        str(missing_stage[:8]) if missing_stage else "30/30 combinations present")

    # ---------------- 14. independent recomputation vs pipeline CSVs --------
    import csv
    t3_path = os.path.join(ROOT, "analysis_outputs", "tables", "T3_Price_Chain.csv")
    t10_path = os.path.join(ROOT, "analysis_outputs", "tables", "T10_Margin_Summary.csv")
    comp_share = np.mean(list(share.values())) if share else None
    if os.path.exists(t3_path) and share:
        with open(t3_path, encoding="utf-8-sig") as fh:
            t3 = {r["Species_code"]: r for r in csv.DictReader(fh)}
        dev = []
        for sp, sh in share.items():
            cell = (t3[sp]["Producer_share_pct"] if sp in t3 else "") or ""
            t3v = float(cell) if str(cell).strip() not in ("", "nan") else None
            if t3v is not None and abs(sh - t3v) > 1.0:
                dev.append((sp, round(sh, 1), t3v))
        rep("PASS" if not dev else "WARN",
            "Recomputed per-species share matches T3 (±1pt)",
            str(dev) if dev else "all reported species match (blank cells = share "
                                 "suppressed on low consumer n, as flagged in T3)")
    if os.path.exists(t10_path) and comp_share:
        with open(t10_path, encoding="utf-8-sig") as fh:
            t10 = {r["Level"]: r for r in csv.DictReader(fh)}
        claimed = float(t10["Producer share of consumer price"]["Margin_pct_consumer"]
                        or 0)
        rep("PASS" if claimed and abs(comp_share - claimed) < 1.5 else "WARN",
            "Recomputed overall share matches T10 claim",
            f"mine={comp_share:.1f}% vs T10={claimed}%")
        for lvl, mkey in [("Aratdar (auction margin, M1/M6)", "aratdar"),
                          ("Bepari/Faria (wholesale margin)", "bepari"),
                          ("Retailer (margin to consumer)", "retail")]:
            v = np.mean(mgn[mkey]) if mgn[mkey] else 0
            c = float(t10[lvl]["Margin_BDT_kg"] or 0)
            rep("PASS" if abs(v - c) < 1.5 else "WARN",
                f"Recomputed {mkey} margin matches T10",
                f"mine={v:.1f} vs T10={c:.1f} BDT/kg")

    # ---------------- 15. K/D flags ----------------
    kd = Counter()
    for rows in (A, B, R, C):
        for r in rows:
            for k, v in r.items():
                if str(v).strip().upper() in ("K", "D"):
                    kd[str(v).strip().upper()] += 1
    rep("PASS", "K/D flags present (realism)", str(dict(kd)))

    spare = []
    for sn in ("Form_A_Aratdar", "Form_B_Bepari_Faria", "Form_R_Khuchra",
               "Form_C_Consumer"):
        ws = wb[sn]
        for r in range(5, ws.max_row + 1):
            rid = ws.cell(row=r, column=2).value
            if rid and ws.cell(row=r, column=4).value in (None, ""):
                # has ID but no Interview_Date -> spare row
                spare.append((sn.split("_")[1], str(rid)))
    rep("PASS" if len(spare) == 24 else "WARN",
        "Spare rows (prefilled ID, unused) documented",
        f"{len(spare)} unused spare IDs, e.g. {spare[:4]}")

    print()
    print("=" * 78)
    print(f"AUDIT SUMMARY:  PASS={n_pass}  WARN={n_warn}  FAIL={n_fail}")
    print("=" * 78)
    wb.close()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=DEFAULT)
    args = ap.parse_args()
    main(args.data)
