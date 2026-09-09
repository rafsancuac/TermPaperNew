#!/usr/bin/env python3
"""Dump structure of the FILLED workbook: headers, row counts, sample rows,
and whether formula cells have cached values (from LibreOffice recalc)."""
import openpyxl

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
F = os.path.join(ROOT, "04_data_filled", "Marine_Fish_Marketing_Data_Entry_Chattogram_Filled.xlsx")

wb = openpyxl.load_workbook(F, data_only=True)  # cached values view
wbf = openpyxl.load_workbook(F, data_only=False)  # formula view

for ws in wb.sheetnames:
    ws, wsf = wb[ws], wbf[ws]
    # find last row with any value in col B..P
    last = 0
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 2000)):
        if any(c.value is not None for c in row[1:16]):
            last = row[0].row
    hdr_row = 4
    headers = []
    for c in range(2, ws.max_column + 1):
        h = ws.cell(row=hdr_row, column=c).value
        if h:
            headers.append((c, str(h)[:28]))
    print(f"\n=== {ws.title} (dims={ws.dimensions}, last_used_row={last}) ===")
    print("HEADERS:", " | ".join(f"{openpyxl.utils.get_column_letter(c)}:{h}" for c, h in headers))
    # sample first data row
    r = 5
    vals = []
    for c in range(2, min(ws.max_column, 40) + 1):
        v = ws.cell(row=r, column=c).value
        fv = wsf.cell(row=r, column=c).value
        tag = ""
        if isinstance(fv, str) and fv.startswith("="):
            tag = f"[F cached={v}]"
        if v is not None or tag:
            vals.append(f"{openpyxl.utils.get_column_letter(c)}={str(v)[:20]}{tag}")
    print("ROW5:", " ; ".join(vals[:22]))
