#!/usr/bin/env python3
"""Inspect the Excel template: sheet names, headers, structure."""
import openpyxl

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(ROOT, "03_data_entry_template", "Marine_Fish_Marketing_Data_Entry (1).xlsx")
wb = openpyxl.load_workbook(PATH, data_only=True)

print("=== SHEET LIST ===")
for ws in wb.worksheets:
    print(f"- {ws.title}: {ws.max_row} rows x {ws.max_column} cols")

print("\n\n=== HEADERS / FIRST ROWS ===")
for ws in wb.worksheets:
    print(f"\n--- {ws.title} ---")
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), 1):
        cells = []
        for c in row:
            s = str(c) if c is not None else ""
            if len(s) > 70:
                s = s[:70] + "..."
            cells.append(s)
        if any(cells):
            print(f"  R{row_idx}: {cells}")
